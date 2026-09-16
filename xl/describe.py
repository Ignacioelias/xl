"""One-call structural map of a workbook: sheets, blocks, headers, links, names."""
import os

from openpyxl.utils import get_column_letter

from .wbcache import resolve_wb

MAX_CELLS_SCAN = 1_500_000
ROW_GAP = 2   # >= this many blank rows ends a band (one spacer row does not)
COL_GAP = 2


def is_formula(v):
    if isinstance(v, str):
        return v.startswith("=")
    return type(v).__name__ in ("ArrayFormula", "DataTableFormula")


def formula_text(v):
    if isinstance(v, str):
        return v
    return getattr(v, "text", "") or ""


def ftext(v):
    """Formula text (starting with '=') for a plain or array/data-table formula, else None.
    Every place that walks formulas must go through this: openpyxl 3.1 returns ArrayFormula /
    DataTableFormula objects for CSE formulas, which `isinstance(v, str)` silently skips."""
    if isinstance(v, str):
        return v if v.startswith("=") else None
    if type(v).__name__ in ("ArrayFormula", "DataTableFormula"):
        t = getattr(v, "text", "") or ""
        return t if t.startswith("=") else ("=" + t if t else None)
    return None


def peek(ws, row, col):
    """The cell at (row, col) or None, WITHOUT creating it. `ws.cell()` / `ws[addr]` insert an
    empty Cell permanently, so a lint pass that probes neighbours grows the cached workbook and
    inflates max_row/max_column for every later reader (measured: 2 cells -> 8 after three reads)."""
    cells = getattr(ws, "_cells", None)
    if cells is None:
        return ws.cell(row, col)
    return cells.get((row, col))


def peek_value(ws, row, col):
    c = peek(ws, row, col)
    return None if c is None else c.value


def grid(ws, addr, max_cells=250_000):
    """Rows of cells (None where empty) for an A1 range, without creating cells. Whole-column
    and whole-row references (A:A, 5:5) are bounded by the sheet's populated extent."""
    from openpyxl.utils.cell import range_boundaries
    c1, r1, c2, r2 = range_boundaries(addr)
    c1, r1 = c1 or 1, r1 or 1
    c2 = c2 or ws.max_column or 1
    r2 = r2 or ws.max_row or 1
    if (c2 - c1 + 1) * (r2 - r1 + 1) > max_cells:
        raise ValueError(f"{addr}: {(c2 - c1 + 1) * (r2 - r1 + 1):,} cells exceeds {max_cells:,}")
    return [[peek(ws, r, c) for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)]


def grid_values(ws, addr, max_cells=250_000):
    return [[None if c is None else c.value for c in row] for row in grid(ws, addr, max_cells)]


def iter_cells(ws):
    """Populated cells only, in (row, col) order. openpyxl keeps them in ws._cells, so this
    never walks a ghost used range (a sheet reporting A1:VWG1315 has ~1,400 real cells)."""
    cells = getattr(ws, "_cells", None)
    if cells is None:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    yield c
        return
    for key in sorted(cells):
        c = cells[key]
        if c.value is not None:
            yield c


def rows_of(ws):
    """Yield (row_index, [populated cells]) in order."""
    cur_r, cur = None, []
    for c in iter_cells(ws):
        if c.row != cur_r:
            if cur:
                yield cur_r, cur
            cur_r, cur = c.row, []
        cur.append(c)
    if cur:
        yield cur_r, cur


def scan_sheet(ws, max_cells=MAX_CELLS_SCAN):
    max_row, max_col = ws.max_row or 0, ws.max_column or 0
    occ = {}
    n_vals = n_form = n_text = n_num = ext = cross = 0
    header = None
    truncated = False
    real_max_col = 0
    for r, cells in rows_of(ws):
        cols, texts = set(), 0
        for c in cells:
            v = c.value
            cols.add(c.column)
            if is_formula(v):
                n_form += 1
                s = formula_text(v)
                if "[" in s:
                    ext += 1
                if "!" in s:
                    cross += 1
            elif isinstance(v, str):
                n_text += 1
                texts += 1
            elif isinstance(v, (int, float)) and not isinstance(v, bool):
                n_num += 1
        occ[r] = cols
        n_vals += len(cells)
        real_max_col = max(real_max_col, max(cols))
        if header is None and texts >= 3 and texts * 2 >= len(cells):
            header = r
        if n_vals >= max_cells:
            truncated = True
            break
    return {"max_row": max_row, "max_col": max_col, "real_max_col": real_max_col, "occ": occ,
            "n_vals": n_vals, "n_form": n_form, "n_text": n_text, "n_num": n_num, "ext": ext,
            "cross": cross, "header": header, "truncated": truncated,
            "scanned_rows": max(occ) if occ else 0}


def _runs(sorted_ints, gap):
    runs, cur, prev = [], [], None
    for x in sorted_ints:
        if prev is not None and x - prev > gap:
            runs.append(cur)
            cur = []
        cur.append(x)
        prev = x
    if cur:
        runs.append(cur)
    return runs


def blocks_from_occ(occ):
    blocks = []
    for band in _runs(sorted(occ), ROW_GAP):
        colset = set()
        for r in band:
            colset |= occ[r]
        for run in _runs(sorted(colset), COL_GAP):
            c1, c2 = run[0], run[-1]
            brows = [r for r in band if any(c1 <= c <= c2 for c in occ[r])]
            if not brows:
                continue
            ncells = sum(1 for r in brows for c in occ[r] if c1 <= c <= c2)
            blocks.append({"r1": brows[0], "r2": brows[-1], "c1": c1, "c2": c2, "ncells": ncells})
    blocks.sort(key=lambda b: -b["ncells"])
    return blocks


def _a1(b):
    return f"{get_column_letter(b['c1'])}{b['r1']}:{get_column_letter(b['c2'])}{b['r2']}"


def _kind(ws, b):
    rows, cols = b["r2"] - b["r1"] + 1, b["c2"] - b["c1"] + 1
    if rows == 1 and cols >= 3:
        return "band"
    if cols == 1:
        return "list"
    first = [peek_value(ws, b["r1"], c) for c in range(b["c1"], b["c2"] + 1)]
    texts = sum(1 for v in first if isinstance(v, str) and not v.startswith("="))
    if cols == 2 and rows >= 2:
        return "kv"
    if texts >= max(2, cols // 2):
        return "table"
    return "group"


def _labels(ws, b, n=6, width=24):
    out = []
    for c in range(b["c1"], b["c2"] + 1):
        v = peek_value(ws, b["r1"], c)
        if v is None:
            continue
        s = str(v)
        if len(s) > width:
            s = s[: width - 1] + "…"
        out.append(s)
        if len(out) >= n:
            break
    col_labels = []
    for r in range(b["r1"] + 1, min(b["r2"], b["r1"] + 40) + 1):
        v = peek_value(ws, r, b["c1"])
        if isinstance(v, str) and not v.startswith("="):
            s = v.strip()
            col_labels.append(s[: width - 1] + "…" if len(s) > width else s)
        if len(col_labels) >= 4:
            break
    return out, col_labels


def describe(path, blocks=8, sample=True):
    wb = resolve_wb(path)
    lines = []
    p = getattr(path, "__fspath__", None) and os.fspath(path) or (path if isinstance(path, str) else None)
    size = os.path.getsize(p) / 1e6 if p and os.path.exists(p) else None
    hidden = [ws.title for ws in wb.worksheets if ws.sheet_state != "visible"]
    calc = wb.calculation
    calc_s = ""
    if calc is not None:
        calc_s = f"calc={calc.calcMode or 'auto'}" + (" iterate=yes" if calc.iterate else "")
    scans = {ws.title: scan_sheet(ws) for ws in wb.worksheets}
    tot_f = sum(s["n_form"] for s in scans.values())
    tot_v = sum(s["n_vals"] for s in scans.values())
    lines.append(f"WORKBOOK {os.path.basename(p) if p else ''}  {f'{size:.1f} MB  ' if size else ''}"
                 f"sheets={len(wb.worksheets)}{f' ({len(hidden)} hidden)' if hidden else ''}  "
                 f"cells={tot_v:,}  formulas={tot_f:,}  vba={'yes' if wb.vba_archive else 'no'}  {calc_s}")
    names = list(wb.defined_names.items()) if hasattr(wb.defined_names, "items") else []
    if names:
        broken = [k for k, v in names if "#REF" in str(v.attr_text)]
        shown = ", ".join(f"{k}={v.attr_text}" for k, v in names[:12])
        tag = f", {len(broken)} broken #REF!" if broken else ""
        lines.append(f"Defined names ({len(names)}{tag}): {shown}{' …' if len(names) > 12 else ''}")
    links = getattr(wb, "_external_links", None) or []
    if links:
        lines.append(f"External links ({len(links)}):")
        for i, ln in enumerate(links, 1):
            tgt = getattr(getattr(ln, "file_link", None), "Target", "?")
            lines.append(f"  [{i}] {tgt}")
    lines.append("SHEETS  (# name | state | used range | values | formulas | header row | cross-sheet refs | external refs)")
    for i, ws in enumerate(wb.worksheets, 1):
        s = scans[ws.title]
        rng = f"A1:{get_column_letter(s['max_col'])}{s['max_row']}" if s["max_col"] else "empty"
        if s["max_col"] and s["real_max_col"] and s["max_col"] > 4 * s["real_max_col"] + 20:
            rng += f" (ghost; real ..{get_column_letter(s['real_max_col'])})"
        lines.append(f"  {i:>2} {ws.title[:32]:<32} {ws.sheet_state:<8} {rng:<28} v={s['n_vals']:<7,} f={s['n_form']:<7,} "
                     f"hdr={s['header'] or '-':<5} x={s['cross']:<6,} ext={s['ext']:,}")
    lines.append("DETAIL")
    for ws in wb.worksheets:
        s = scans[ws.title]
        if not s["n_vals"]:
            continue
        extras = []
        if ws.freeze_panes:
            extras.append(f"freeze {ws.freeze_panes}")
        m = len(ws.merged_cells.ranges)
        if m:
            extras.append(f"merged {m}")
        if ws.tables:
            extras.append("tables " + ",".join(f"{t.name}[{t.ref}]" for t in ws.tables.values()))
        hr = sum(1 for d in ws.row_dimensions.values() if d.hidden)
        hc = sum(1 for d in ws.column_dimensions.values() if d.hidden)
        if hr or hc:
            extras.append(f"hidden r{hr}/c{hc}")
        ch = len(getattr(ws, "_charts", []))
        if ch:
            extras.append(f"charts {ch}")
        if s["truncated"]:
            extras.append(f"scan capped at {MAX_CELLS_SCAN:,} cells (row {s['scanned_rows']})")
        lines.append(f"- {ws.title}  [{', '.join(extras)}]" if extras else f"- {ws.title}")
        bl = blocks_from_occ(s["occ"])
        for b in bl[:blocks]:
            rows, cols = b["r2"] - b["r1"] + 1, b["c2"] - b["c1"] + 1
            kind = _kind(ws, b) if sample else ""
            top, left = _labels(ws, b) if sample else ([], [])
            desc = f"    {_a1(b):<14} {rows:>5}x{cols:<4} {b['ncells']:>7,} cells  {kind:<6}"
            if top:
                desc += "  top: " + " | ".join(top)
            if left:
                desc += "  rows: " + " | ".join(left)
            lines.append(desc)
        if len(bl) > blocks:
            lines.append(f"    … {len(bl) - blocks} smaller blocks")
    return "\n".join(lines)

"""Formula analysis on the cached openpyxl workbook: references, precedents, dependents,
multi-level trace, label lookup, find with row context, find-and-replace.

Reference parsing uses openpyxl's Tokenizer, so LOG10 / ATAN2 never read as cell refs.
"""
import functools
import re
from array import array

from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from .describe import ftext, grid_values, iter_cells, peek_value, rows_of
from .wbcache import get_sheet, resolve_wb, split_ref

_ADDR = re.compile(r"^\$?[A-Za-z]{1,3}\$?\d+(?::\$?[A-Za-z]{1,3}\$?\d+)?$|^\$?[A-Za-z]{1,3}:\$?[A-Za-z]{1,3}$|^\$?\d+:\$?\d+$")
_INDEX_ATTR = "_xl_ref_index"


@functools.lru_cache(maxsize=300_000)
def tokens_of(formula):
    """Tokenizer output for a formula string, memoised: the lint rules tokenise the same formula
    up to eight times across their walks, and translated copies repeat verbatim across rows.
    Returns () when the tokenizer rejects the text."""
    try:
        return tuple(Tokenizer(formula).items)
    except Exception:
        return ()


def refs_in(formula, default_sheet, wb=None):
    """Yield (sheet, a1_range, is_external) for every range operand in a formula (plain text or
    an openpyxl ArrayFormula / DataTableFormula). Defined names resolve through the workbook."""
    formula = ftext(formula)
    if formula is None:
        return
    tokens = tokens_of(formula)
    for t in tokens:
        if t.type != "OPERAND" or t.subtype != "RANGE":
            continue
        v = t.value
        external = v.startswith("[") or "]" in v.split("!")[0]
        sheet, addr = default_sheet, v
        if "!" in v:
            sheet, addr = v.rsplit("!", 1)
            if sheet.startswith("'") and sheet.endswith("'"):
                sheet = sheet[1:-1].replace("''", "'")
        addr = addr.replace("$", "")
        if addr.endswith("#"):
            addr = addr[:-1]
        if _ADDR.match(addr):
            yield sheet, addr.upper(), external
        elif wb is not None and not external:
            dn = wb.defined_names.get(addr) if hasattr(wb.defined_names, "get") else None
            if dn is not None and "#REF" not in str(dn.attr_text):
                try:
                    for s, r in dn.destinations:
                        yield s, r.replace("$", "").upper(), False
                except Exception:
                    continue


def _bounds(addr, ws):
    c1, r1, c2, r2 = range_boundaries(addr)
    return (c1 or 1, r1 or 1, c2 or ws.max_column or 1, r2 or ws.max_row or 1)


def precedents(path, ref, values=True, max_cells=12):
    wf = resolve_wb(path)
    wv = resolve_wb(path, values=True) if values else wf
    sh, addr = split_ref(ref)
    ws = get_sheet(wf, sh)
    c1, r1, _c2, _r2 = range_boundaries(addr)
    f = peek_value(ws, r1 or 1, c1 or 1)
    rows = [["formula", f"{ws.title}!{addr}", ftext(f) or f]]
    seen = set()
    for s, a, ext in refs_in(f, ws.title, wf):
        if (s, a) in seen:
            continue
        seen.add((s, a))
        if ext:
            rows.append(["external", f"{s}!{a}", "(other workbook; cached only)"])
            continue
        try:
            wsv = get_sheet(wv, s)
            vals = grid_values(wsv, a)
        except Exception as e:  # noqa: BLE001
            rows.append(["ref", f"{s}!{a}", f"<{e}>"])
            continue
        if ":" not in a:
            rows.append(["ref", f"{s}!{a}", vals[0][0]])
        else:
            flat = [v for row in vals for v in row]
            rows.append(["range", f"{s}!{a}", f"{len(flat)} cells: {flat[:max_cells]}{' …' if len(flat) > max_cells else ''}"])
    return rows


def _index(wb):
    """Per referenced sheet: parallel int arrays (c1, r1, c2, r2, src_sheet_idx, src_row, src_col).
    Stored on the Workbook object itself, so a reloaded workbook (after calc --save, save_as or
    replace) starts with no index and the old one dies with the old object. A module-level dict
    keyed by id(wb) served the pre-edit index to a reloaded workbook at the same address (28 of 30
    reloads in the probe) and leaked seven int arrays per sheet per reload."""
    cached = getattr(wb, _INDEX_ATTR, None)
    if cached is not None:
        return cached
    idx = {}
    names = [ws.title for ws in wb.worksheets]
    low = {n.lower(): n for n in names}
    for si, ws in enumerate(wb.worksheets):
        for c in iter_cells(ws):
            v = ftext(c.value)
            if v is None:
                continue
            for s, a, ext in refs_in(v, ws.title, wb):
                if ext:
                    continue
                s_norm = low.get(s.lower())
                if s_norm is None:
                    continue
                try:
                    b = _bounds(a, wb[s_norm])
                except Exception:
                    continue
                e = idx.setdefault(s_norm, [array("i") for _ in range(7)])
                for arr, val in zip(e, (*b, si, c.row, c.column)):
                    arr.append(int(val))
    setattr(wb, _INDEX_ATTR, (idx, names))
    return getattr(wb, _INDEX_ATTR)


def dependents(path, ref, limit=200):
    """Cells whose formulas read `ref` (direct dependents), across all sheets."""
    wb = resolve_wb(path)
    sh, addr = split_ref(ref)
    ws = get_sheet(wb, sh)
    c1, r1, c2, r2 = _bounds(addr, ws)
    idx, names = _index(wb)
    e = idx.get(ws.title)
    out = []
    if not e:
        return out
    C1, R1, C2, R2, SI, SR, SC = e
    for i in range(len(C1)):
        if C1[i] <= c2 and C2[i] >= c1 and R1[i] <= r2 and R2[i] >= r1:
            out.append((names[SI[i]], f"{get_column_letter(SC[i])}{SR[i]}"))
            if len(out) >= limit:
                break
    return out


def trace(path, ref, depth=3, direction="inputs", max_per_level=25):
    """Multi-level trace. direction='inputs' walks precedents until constants (traceToInputs);
    'outputs' walks dependents (traceToOutputs). Returns rows [level, cell, kind, value/formula]."""
    wf = resolve_wb(path)
    wv = resolve_wb(path, values=True)
    sh, addr = split_ref(ref)
    start = get_sheet(wf, sh).title
    rows, seen, frontier = [], set(), [(start, addr.upper())]
    for level in range(depth + 1):
        nxt = []
        for s, a in frontier[:max_per_level]:
            if (s, a) in seen:
                continue
            seen.add((s, a))
            ws = get_sheet(wf, s)
            try:
                ac, ar, _, _ = range_boundaries(a)
                f = peek_value(ws, ar, ac)
            except Exception:
                continue
            val = None
            try:
                val = peek_value(get_sheet(wv, s), ar, ac)
            except Exception:
                pass
            f = ftext(f) or f
            is_f = isinstance(f, str) and f.startswith("=")
            rows.append([level, f"{s}!{a}", "formula" if is_f else "input", f if is_f else val, val if is_f else None])
            if direction == "inputs":
                if is_f:
                    for ps, pa, ext in refs_in(f, s, wf):
                        if ext:
                            continue
                        if ":" in pa:
                            c1, r1, c2, r2 = _bounds(pa, get_sheet(wf, ps))
                            if (c2 - c1 + 1) * (r2 - r1 + 1) > 12:
                                rows.append([level + 1, f"{ps}!{pa}", "range", f"{(c2 - c1 + 1) * (r2 - r1 + 1)} cells", None])
                                continue
                            for rr in range(r1, r2 + 1):
                                for cc in range(c1, c2 + 1):
                                    nxt.append((ps, f"{get_column_letter(cc)}{rr}"))
                        else:
                            nxt.append((ps, pa))
            else:
                for ds, da in dependents(path, f"{s}!{a}", limit=max_per_level):
                    nxt.append((ds, da))
        if not nxt:
            break
        frontier = nxt
    return rows


def lookup(path, table, row_label, col_label=None, values=True, sheet=None):
    """Find a cell by human labels: `table` is a sheet name or 'Sheet!A1:H40'. row_label is
    matched (regex, case-insensitive) in the first 6 columns of the block; col_label in the
    first 12 rows. Returns [address, value, row_label_cell, col_label_cell]."""
    wb = resolve_wb(path, values=values)
    sh, rng = split_ref(table) if "!" in table else (table, None)
    ws = get_sheet(wb, sh)
    if rng:
        c1, r1, c2, r2 = _bounds(rng, ws)
    else:
        c1, r1, c2, r2 = 1, 1, ws.max_column, ws.max_row

    def _rx(label):
        # labels are regexes ("EBITDA|Op\. profit"), but an unbalanced "Revenue (net" must not
        # blow up with re.error: fall back to a literal match
        try:
            return re.compile(str(label), re.I)
        except re.error:
            return re.compile(re.escape(str(label)), re.I)

    rrx = _rx(row_label)
    crx = _rx(col_label) if col_label not in (None, "") else None
    row_hits, col_hit = [], None
    for r in range(r1, r2 + 1):
        for c in range(c1, min(c2, c1 + 5) + 1):
            v = peek_value(ws, r, c)
            if isinstance(v, str) and rrx.search(v):
                row_hits.append((r, c))
                break
    row_hit = row_hits[0] if row_hits else None
    if crx:
        for r in range(r1, min(r2, r1 + 11) + 1):
            for c in range(c1, c2 + 1):
                v = peek_value(ws, r, c)
                if v is not None and crx.search(str(v)):
                    col_hit = (r, c)
                    break
            if col_hit:
                break
    if not row_hit:
        raise KeyError(f"row label /{row_label}/ not found in {ws.title}")
    if crx and not col_hit:
        raise KeyError(f"column label /{col_label}/ not found in {ws.title}")
    # A label can appear several times (section header row, then the numeric row). Take the
    # first hit whose target cell holds a value; fall back to the first hit if none does.
    def _target(hit):
        r = hit[0]
        if col_hit:
            return r, col_hit[1]
        return r, next((cc for cc in range(hit[1] + 1, c2 + 1)
                        if isinstance(peek_value(ws, r, cc), (int, float))), hit[1] + 1)

    r, c = _target(row_hit)
    if peek_value(ws, r, c) is None:
        for hit in row_hits[1:]:
            rr, cc = _target(hit)
            if peek_value(ws, rr, cc) is not None:
                row_hit, r, c = hit, rr, cc
                break
    a = f"{get_column_letter(c)}{r}"
    return [f"{ws.title}!{a}", peek_value(ws, r, c),
            f"{get_column_letter(row_hit[1])}{row_hit[0]}",
            f"{get_column_letter(col_hit[1])}{col_hit[0]}" if col_hit else None]


def find_rows(path, pattern, sheet=None, context=0, limit=20, values=False, flags=re.I):
    """Rows where any cell matches; returns the populated cells of the row (and +-context rows)."""
    wb = resolve_wb(path, values=values)
    rx = re.compile(pattern, flags)
    out = []
    targets = [get_sheet(wb, sheet)] if sheet else wb.worksheets
    for ws in targets:
        rows = list(rows_of(ws))
        by_r = {r: cells for r, cells in rows}
        for i, (r, cells) in enumerate(rows):
            if any(rx.search(str(c.value)) for c in cells):
                for rr in range(r - context, r + context + 1):
                    if rr in by_r:
                        out.append([f"{ws.title}!{rr}", " | ".join(f"{c.coordinate}={c.value}" for c in by_r[rr][:14])])
                out.append(["", ""])
                if sum(1 for o in out if o[0] == "") >= limit:
                    return out
    return out


def find_and_replace(path, search, replace, in_formulas=False, sheet=None, whole_cell=False,
                     match_case=False, dry_run=True, force=False):
    """Replace text in cell values (and formulas if asked). dry_run=True only reports.
    A write re-serialises the file through openpyxl, so it is refused on any original that is
    not a work copy (sync root or not) unless force=True, and a backup goes to the work dir first."""
    from .paths import guard_write
    if not dry_run:
        guard_write(path, force=force)
    wb = resolve_wb(path)
    flags = 0 if match_case else re.I
    rx = re.compile(re.escape(search) if not whole_cell else f"^{re.escape(search)}$", flags)
    changes = []
    targets = [get_sheet(wb, sheet)] if sheet else wb.worksheets
    for ws in targets:
        for c in iter_cells(ws):
            v = c.value
            if not isinstance(v, str):
                continue
            if v.startswith("=") and not in_formulas:
                continue
            nv, n = rx.subn(replace, v)
            if n:
                changes.append([f"{ws.title}!{c.coordinate}", v[:60], nv[:60]])
                if not dry_run:
                    c.value = nv
    if not dry_run and changes:
        from .paths import backup_to_work
        backup_to_work(path)
        wb.save(path)
        from . import wbcache
        wbcache.drop(path)
    return changes

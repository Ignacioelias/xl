"""The namespace preloaded into every kernel session. `xl help` prints API."""
import json
import math
import os
import re
import shutil
import time
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from . import describe as _describe_mod
from .describe import grid, iter_cells
from . import excelcom as _xl
from . import lint as _lint_mod
from . import wbcache
from . import analysis as _an
from . import cached as _cached
from . import pptcom as _pp
from .paths import WORK, under_sync_root
from .wbcache import get_sheet, resolve_wb, split_ref

API = """xl kernel API (state persists across calls; `_` is the last result)

Workbooks (openpyxl, cached by path+mtime; nothing here recalculates)
  wb(path, values=False, reload=False)   formulas as text; values=True -> cached results at last save
  sheets(path)                            index, name, state, used range
  describe(path, blocks=8)                one-call structural map (sheets, blocks, headers, links, names)
  find(path, pattern, values=False, sheet=None, limit=50)   regex over cell text (values=True searches cached results)
  read(path, "Sheet!A1:H30", values=True) -> Table of rows; values=False gives formulas
  read_tsv(path, "Sheet!A1:H30", values=True)   dense TSV with column letters and row numbers (broad inspection)
  count(path, pattern, values=False, sheet=None)  number of matching cells, no output cap
  precedents(path, "Sheet!C12")           cells/ranges a formula reads, with their cached values
  dependents(path, "Sheet!C12")           cells whose formulas read it (all sheets)
  trace(path, "Sheet!C12", depth=3, direction="inputs"|"outputs")   multi-level walk
  lookup(path, "Sheet" or "Sheet!A1:H40", row_label, col_label=None)  cell by human labels
  find_rows(path, pattern, sheet=None, context=0)   matching rows with neighbours (context=N adds N rows each side)
  replace(path, search, repl, in_formulas=False, dry_run=True, force=False)   find-and-replace; writes work copies only unless force=True (backup taken)
  work_copy(path)                         copy into the work dir and return the new path (never edit originals)
  save_as(wb_obj, path, force=False)      refuses to write under a sync root unless force=True; drops the cache entry
  cache_info(), drop_cache(path=None)

Engine (dedicated hidden Excel instance: numbers come from Excel; check-open is the repair-prompt test)
  calc(path, save=False, verify=False, force=False)   full rebuild; every error cell; verify adds changed cached values
                                          (volatile NOW/TODAY/RAND excluded); save= only on a work copy unless force=True, backup taken
  scenario(path, {"Sheet!B3": 0.05}, ["Sheet!B40"])   what-if without saving: base, scenario, delta
  sweep(path, "Sheet!B3", [0.03, 0.05, 0.07], ["Sheet!B40"])   one-variable sweep
  render(path, "Sheet!A1:H30", out=None, diff=None)  dict {png, bytes, range, method}; diff=baseline.png adds a pixel diff
  pptx_lint(path, slides=None)   off-slide, text overflow, occluded text, tiny fonts, empty placeholders/charts
  pptx_render(path, slide=None, out_dir=None)   slide PNGs via PowerPoint (opens a COPY, never quits the app)
  check_open(path)           open in a throwaway Excel with alerts ON; blocked => repair prompt => NOT verified

Package level (no Excel; think-cell feeds and other blocks that must read right without a recalculation)
  inject_cached(path, "Sheet", {"C5": 1.5, "D5": "FY25"}, out=None, force=False)
                             write cached values into existing formula cells; one <v> per cell asserted on read-back;
                             in place only on a work copy unless force=True (backup taken); warns if calcPr would recalc on open
  calcpr(path, set=None, out=None, force=False)   read <calcPr>, or set={"calcMode": "manual", "fullCalcOnLoad": None}

Hard rules
  1. After any write: calc(path) and inspect ['errors'] before reporting done.
  2. Numbers come from the engine, not from Python arithmetic. Write the formula, calc, read back.
  3. Cite Sheet!Address for every figure you report.
  4. Read rows through the file; never hide/unhide, never merge.
"""


class Table(list):
    """List of rows that prints as an aligned text table (capped)."""
    max_rows = 200
    max_width = 40

    def __repr__(self):
        if not self:
            return "(empty)"
        rows = [[("" if v is None else str(v)) for v in r] for r in self[: self.max_rows]]
        rows = [[(s if len(s) <= self.max_width else s[: self.max_width - 1] + "…") for s in r] for r in rows]
        ncol = max(len(r) for r in rows)
        widths = [0] * ncol
        for r in rows:
            for i, s in enumerate(r):
                widths[i] = max(widths[i], len(s))
        lines = ["  ".join(s.ljust(widths[i]) for i, s in enumerate(r)).rstrip() for r in rows]
        if len(self) > self.max_rows:
            lines.append(f"… {len(self) - self.max_rows} more rows")
        return "\n".join(lines)

    def tsv(self, max_rows=None):
        """Tab-separated text: the cheapest faithful view of a table for an agent."""
        rows = self if max_rows is None else self[:max_rows]
        return "\n".join("\t".join("" if v is None else str(v) for v in r) for r in rows)


def _norm(path):
    if isinstance(path, str):
        m = re.match(r"^/([A-Za-z])/(.*)$", path)
        if m:
            path = f"{m.group(1).upper()}:/{m.group(2)}"
    return path


def build_namespace(session):
    def wb(path, values=False, reload=False):
        return wbcache.open_wb(_norm(path), values=values, reload=reload)

    def sheets(path):
        w = resolve_wb(_norm(path))
        t = Table()
        for i, ws in enumerate(w.worksheets, 1):
            t.append([i, ws.title, ws.sheet_state, ws.dimensions])
        return t

    def describe(path, blocks=8):
        return _describe_mod.describe(_norm(path), blocks=blocks)

    def find(path, pattern, values=False, sheet=None, limit=50, flags=re.I):
        w = resolve_wb(_norm(path), values=values)
        rx = re.compile(pattern, flags)
        t = Table()
        targets = [get_sheet(w, sheet)] if sheet else w.worksheets
        for ws in targets:
            for c in iter_cells(ws):
                v = c.value
                s = v if isinstance(v, str) else str(v)
                if rx.search(s):
                    t.append([f"{ws.title}!{c.coordinate}", s])
                    if len(t) >= limit:
                        return t
        return t

    def read(path, ref, values=True):
        w = resolve_wb(_norm(path), values=values)
        sh, rng = split_ref(ref)
        ws = get_sheet(w, sh)
        t = Table()
        for row in grid(ws, rng):  # non-creating read: ws[rng] would insert empty cells into the cache
            t.append([None if c is None else c.value for c in row])
        return t

    def read_tsv(path, ref, values=True):
        """Range as TSV with a column-letter header and row numbers, so every value keeps its
        address. Dense and token-cheap; use it for broad inspection, `read` for a few cells."""
        from openpyxl.utils.cell import range_boundaries
        w = resolve_wb(_norm(path), values=values)
        sh, rng = split_ref(ref)
        ws = get_sheet(w, sh)
        c1, r1, c2, r2 = range_boundaries(rng)
        c1, r1 = c1 or 1, r1 or 1
        c2, r2 = c2 or ws.max_column or 1, r2 or ws.max_row or 1
        rows = grid(ws, rng)
        head = [f"{ws.title}!"] + [get_column_letter(c) for c in range(c1, c2 + 1)]
        lines = ["\t".join(head)]
        for i, row in enumerate(rows):
            lines.append("\t".join([str(r1 + i)] + ["" if c is None or c.value is None else str(c.value) for c in row]))
        return "\n".join(lines)

    def count(path, pattern, values=False, sheet=None, flags=re.I):
        """Number of cells whose text (or cached value with values=True) matches; no output cap.
        Streams over populated cells; never materialises the hits (a 636k-formula model can match
        hundreds of thousands of times)."""
        w = resolve_wb(_norm(path), values=values)
        rx = re.compile(pattern, flags)
        targets = [get_sheet(w, sheet)] if sheet else w.worksheets
        n = 0
        for ws in targets:
            for c in iter_cells(ws):
                v = c.value
                if rx.search(v if isinstance(v, str) else str(v)):
                    n += 1
        return n

    def precedents(path, ref):
        return Table(_an.precedents(_norm(path), ref))

    def dependents(path, ref, limit=200):
        return Table([[f"{s}!{a}"] for s, a in _an.dependents(_norm(path), ref, limit)])

    def trace(path, ref, depth=3, direction="inputs"):
        return Table(_an.trace(_norm(path), ref, depth=depth, direction=direction))

    def lookup(path, table, row_label, col_label=None, values=True):
        return _an.lookup(_norm(path), table, row_label, col_label, values=values)

    def find_rows(path, pattern, sheet=None, context=0, limit=20, values=False):
        return Table(_an.find_rows(_norm(path), pattern, sheet, context, limit, values))

    def replace(path, search, repl, in_formulas=False, sheet=None, whole_cell=False, dry_run=True, force=False):
        return Table(_an.find_and_replace(_norm(path), search, repl, in_formulas, sheet, whole_cell,
                                          dry_run=dry_run, force=force))

    def scenario(path, inputs, outputs):
        res = _xl.scenario(_norm(path), inputs, outputs)
        return Table([["output", "base", "scenario", "delta"]] + res["rows"])

    def sweep(path, input_ref, values, outputs):
        res = _xl.sweep(_norm(path), input_ref, values, outputs)
        return Table([[input_ref] + list(outputs)] + res["rows"])

    def pptx_lint(path, slides=None, as_dict=False):
        res = _pp.lint(_norm(path), slides=slides)
        return res if as_dict else _pp.fmt_lint(res)

    def pptx_render(path, slide=None, out_dir=None):
        return _pp.render(_norm(path), slide=slide, out_dir=_norm(out_dir) if out_dir else None)

    def lint(path, quick=False, as_dict=False):
        return _lint_mod.lint(_norm(path), quick=quick, as_dict=as_dict)

    def calc(path, save=False, verify=False, force=False):
        path = _norm(path)
        res = _xl.calc(path, save=save, verify=verify, force=force)
        if save:
            wbcache.drop(path)
        return res

    def render(path, ref=None, out=None, diff=None):
        return _xl.render(_norm(path), ref, _norm(out) if out else None, diff=_norm(diff) if diff else None)

    def check_open(path, timeout=90):
        return _xl.check_open(_norm(path), timeout=timeout)

    def inject_cached(path, sheet, values, out=None, force=False):
        path = _norm(path)
        res = _cached.inject(path, sheet, values, out=_norm(out) if out else None, force=force)
        wbcache.drop(res["file"])
        return res

    def calcpr(path, set=None, out=None, force=False):  # noqa: A002
        path = _norm(path)
        res = _cached.calcpr(path, set=set, out=_norm(out) if out else None, force=force)
        wbcache.drop(res["file"])
        return res

    def work_copy(path):
        src = Path(_norm(path))
        dst = WORK / f"{time.strftime('%Y%m%d_%H%M%S')}_{src.name}"
        shutil.copyfile(src, dst)
        return str(dst)

    def save_as(wb_obj, path, force=False):
        path = _norm(path)
        if under_sync_root(path) and not force:
            raise PermissionError(f"{path} is under a OneDrive/Teams sync root. Build in the work dir and deliver through the gate, or pass force=True.")
        wb_obj.save(path)
        wbcache.drop(path)
        return path

    def fmt_calc(res, show=25):
        errs = res["errors"]
        by = {}
        for e in errs:
            by.setdefault((e["sheet"], e["error"]), []).append(e)
        lines = [f"CALC {os.path.basename(res['file'])}  sheets={res['sheets']}  "
                 f"error_cells={res['error_count']}{'+' if res['capped'] else ''}  "
                 f"{'SAVED' if res['saved'] else 'not saved'}  {res['secs']}s"
                 + (f"  changed_values={res['changed_count']}{'+' if res['changed_count'] >= 200 else ''}" if 'changed_count' in res else "")
                 + (f"  (volatile skipped: {res['volatile_skipped']})" if res.get('volatile_skipped') else "")
                 + (f"\n  backup: {res['backup']}" if res.get('backup') else "")]
        for ch in (res.get("changed") or [])[:10]:
            lines.append(f"  changed  {ch['sheet']}!{ch['cell']}: {ch['before']!r} -> {ch['after']!r}")
        for (sheet, err), lst in sorted(by.items(), key=lambda kv: -len(kv[1])):
            cells = ", ".join(e["cell"] for e in lst[:8]) + (" …" if len(lst) > 8 else "")
            lines.append(f"  {err:<8} {sheet:<28} {len(lst):>4} cells  {cells}")
            f0 = next((e["formula"] for e in lst if e["formula"]), None)
            if f0:
                lines.append(f"           e.g. {lst[0]['cell']}: {f0[:100]}")
            if len(lines) > show:
                lines.append("  …")
                break
        return "\n".join(lines)

    def fmt_render(res):
        line = f"RENDER {res['range']} -> {res['png']}  ({res['bytes']:,} bytes, {res['method']})"
        if "diff_png" in res:
            line += f"\n  DIFF {res['changed_pct']}% pixels changed -> {res['diff_png']}" + (f"  ({res['note']})" if res.get("note") else "")
        return line

    def fmt_check(res):
        dd = res.get("dismissed_dialogs") or []
        tail = f"  (auto-closed add-in dialogs: {', '.join(sorted(set(dd)))})" if dd else ""
        if res.get("opened"):
            return (f"CHECK-OPEN ok: {res['name']} opened in {res['secs']}s, sheets={res['sheets']}, "
                    f"read_only={res['read_only']}, no repair prompt{tail}")
        return f"CHECK-OPEN FAILED: {res.get('msg')}{tail}"

    def help():  # noqa: A001 - intentional shadow inside the kernel
        print(API)

    def _shutdown():
        _xl.quit_app()

    ns = {
        "__name__": "__xl__", "__builtins__": __builtins__,
        "openpyxl": openpyxl, "get_column_letter": get_column_letter,
        "column_index_from_string": column_index_from_string,
        "re": re, "json": json, "math": math, "os": os, "Path": Path, "time": time,
        "Table": Table, "WORK": WORK,
        "wb": wb, "open_wb": wb, "sheets": sheets, "describe": describe, "find": find, "read": read,
        "read_tsv": read_tsv, "count": count,
        "precedents": precedents, "dependents": dependents, "trace": trace, "lookup": lookup,
        "find_rows": find_rows, "replace": replace, "scenario": scenario, "sweep": sweep,
        "pptx_lint": pptx_lint, "pptx_render": pptx_render,
        "lint": lint, "calc": calc, "render": render, "check_open": check_open,
        "inject_cached": inject_cached, "calcpr": calcpr,
        "fmt_inject": _cached.fmt_inject, "fmt_calcpr": _cached.fmt_calcpr,
        "work_copy": work_copy, "save_as": save_as, "cache_info": wbcache.cache_info,
        "drop_cache": wbcache.drop, "help": help, "API": API, "SESSION": session,
        "fmt_calc": fmt_calc, "fmt_render": fmt_render, "fmt_check": fmt_check,
        "_shutdown": _shutdown, "_excel_pid": _xl.excel_pid,
    }
    return ns

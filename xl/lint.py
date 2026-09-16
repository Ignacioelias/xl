"""openpyxl-only lint rules. Fast enough for a hook in --quick mode (single load)."""
import re

from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

from . import lint_rules as _LR
from . import lint_rules2 as _LR2
from .describe import ftext, iter_cells, peek_value, rows_of
from .wbcache import open_wb

ERR_TOKENS = {"#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}
NUM_TEXT = re.compile(r"^\s*-?\(?\d{1,3}([.,]\d{3})+([.,]\d+)?\)?\s*%?$|^\s*-?\d+[.,]\d+\s*%?$|^\s*-?\d{4,}\s*$")
LEVEL_ORDER = {"error": 0, "warn": 1, "info": 2}

INLINE_IDS = {"external-link", "defined-name-broken", "iterative-calc", "merged-cells", "hidden",
              "error-constant", "inconsistent-formula", "hardcode-in-formula-row", "number-as-text",
              "cached-error"}
# Ids each registered rule function emits, so a function whose ids are all switched off is not run.
EMITS = {
    "rule_sum_misses_adjacent": {"sum-misses-adjacent"},
    "rule_double_counting": {"double-counting"},
    "rule_mixed_percent": {"mixed-percent"},
    "rule_chart_series": {"chart-series-broken", "chart-series-empty", "chart-series-errors"},
    "rule_object_covers_cells": {"object-covers-cells"},
    "rule_row_format_inconsistent": {"row-format-inconsistent"},
    "rule_external_link_roots": {"external-link-absolute"},
    "rule_empty_cell_coercion": {"empty-cell-coercion"},
    "rule_column_clipping": {"column-clipping"},
    "rule_data_validation": {"data-validation-breach"},
    "rule_lookups": {"unsorted-lookup", "duplicate-lookup-keys"},
    "rule_aggregate_over_text": {"aggregate-over-text"},
    "rule_currency_mixes": {"mixed-currency", "currency-date-mix"},
    "rule_broadcast_surprise": {"broadcast-surprise"},
    "rule_row_height_clipping": {"row-height-clipping"},
    "rule_chart_geometry": {"chart-invisible", "chart-legend-crowded", "chart-axis-labels-crowded"},
}
ALL_IDS = frozenset(INLINE_IDS.union(*EMITS.values()))  # 32 implemented

# Checks switched off on 2026-09-16 (taken off the read-me): never run, never reported. The code
# stays in lint_rules*.py; delete an id here to bring its check back.
OFF_IDS = frozenset({
    "mixed-percent", "chart-series-empty", "chart-series-errors", "object-covers-cells",
    "column-clipping", "data-validation-breach", "row-format-inconsistent", "unsorted-lookup",
    "duplicate-lookup-keys", "broadcast-surprise", "row-height-clipping", "chart-invisible",
    "chart-legend-crowded", "chart-axis-labels-crowded",
})

# Rule ids this linter emits (18). The fixture check asserts each fires and no OFF id does.
RULE_IDS = sorted(ALL_IDS - OFF_IDS)

# The single registry. A rule opts out of --quick with `rule.quick = False` where it is defined.
# A function missing from EMITS still runs; the fixture check flags the ids it emits.
FULL_RULES = [r for r in list(_LR.RULES) + list(_LR2.RULES) if EMITS.get(r.__name__, {None}) - OFF_IDS]
QUICK_RULES = [r for r in FULL_RULES if getattr(r, "quick", True)]


def lint(path, quick=False, max_per_rule=40, as_dict=False):
    wb = open_wb(path)
    findings = []
    counts = {}

    def add(level, rule, where, msg):
        if rule in OFF_IDS:  # rule_chart_series still runs for chart-series-broken
            return
        counts[rule] = counts.get(rule, 0) + 1
        if counts[rule] <= max_per_rule:
            findings.append({"level": level, "rule": rule, "where": where, "msg": msg})

    links = getattr(wb, "_external_links", None) or []
    for i, ln in enumerate(links, 1):
        tgt = getattr(getattr(ln, "file_link", None), "Target", "?")
        add("info", "external-link", f"[{i}]", str(tgt))
    names = list(wb.defined_names.items()) if hasattr(wb.defined_names, "items") else []
    for k, v in names:
        if "#REF" in str(v.attr_text):
            add("error", "defined-name-broken", k, f"{v.attr_text} (named range points at deleted cells)")
    calc = wb.calculation
    if calc is not None and calc.iterate:
        add("info", "iterative-calc", "workbook", "iterate=True: circular references are tolerated, so a broken loop shows as a number, not an error")

    merged_by_sheet, hidden_by_sheet = [], []
    for ws in wb.worksheets:
        merged = list(ws.merged_cells.ranges)
        if merged:
            merged_by_sheet.append((ws.title, len(merged), ", ".join(str(m) for m in merged[:3])))
        hr = sum(1 for d in ws.row_dimensions.values() if d.hidden)
        hc = sum(1 for d in ws.column_dimensions.values() if d.hidden)
        if hr or hc:
            hidden_by_sheet.append((ws.title, hr, hc))
    if merged_by_sheet:
        total = sum(n for _, n, _ in merged_by_sheet)
        ex = "; ".join(f"{t}: {n} ({e})" for t, n, e in merged_by_sheet[:4])
        add("warn", "merged-cells", f"{len(merged_by_sheet)} sheets",
            f"{total} merged ranges (they break sort, filter and copy; use Center Across Selection instead). {ex}{' …' if len(merged_by_sheet) > 4 else ''}")
    if hidden_by_sheet:
        ex = ", ".join(f"{t} r{hr}/c{hc}" for t, hr, hc in hidden_by_sheet[:6])
        add("info", "hidden", f"{len(hidden_by_sheet)} sheets",
            f"hidden rows/cols present: {ex}{' …' if len(hidden_by_sheet) > 6 else ''} (do not unhide; read them through the file)")
    for ws in wb.worksheets:
        for _r, rcells in rows_of(ws):
            cells = [(c.column, c) for c in rcells]
            for col, c in cells:
                v = c.value
                if isinstance(v, str) and v.strip() in ERR_TOKENS:
                    add("error", "error-constant", f"{ws.title}!{c.coordinate}", f"{v.strip()} typed as a constant")
            formulas = [(col, c, ftext(c.value)) for col, c in cells if ftext(c.value)]
            numerics = [(col, c) for col, c in cells
                        if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)]
            years = sum(1 for _, c in numerics if isinstance(c.value, int) and 1990 <= c.value <= 2100)
            is_year_row = numerics and years >= max(3, 0.7 * len(numerics))
            # inconsistent formulas: a break AFTER a run of >= 2 consistent copies (the first
            # forecast column legitimately differs from the rest; a totals column is reported)
            run = 0
            for (c1, a, fa), (c2, b, fb) in zip(formulas, formulas[1:]):
                if c2 - c1 != 1:
                    run = 0
                    continue
                try:
                    expected = Translator(fa, origin=a.coordinate).translate_formula(b.coordinate)
                except Exception:
                    run = 0
                    continue
                if expected == fb:
                    run += 1
                    continue
                if run >= 2:
                    # cohort "staircase": the break matches the cell diagonally up-left -> by design
                    diag = ftext(peek_value(ws, b.row - 1, c2 - 1)) if b.row > 1 and c2 > 1 else None
                    staircase = False
                    if diag:
                        staircase = diag == fb  # same text down the diagonal (cohort anchors)
                        if not staircase:
                            try:
                                origin = f"{get_column_letter(c2 - 1)}{b.row - 1}"
                                staircase = Translator(diag, origin=origin).translate_formula(b.coordinate) == fb
                            except Exception:
                                staircase = False
                    if not staircase:
                        add("warn", "inconsistent-formula", f"{ws.title}!{b.coordinate}",
                            f"breaks the pattern held since {get_column_letter(c2 - run - 1)}: {fb[:90]}")
                        break
                run = 0
            # hardcoded numbers with a formula on BOTH sides (an opening 0 in the first period is fine)
            if len(formulas) >= 4 and not is_year_row:
                fcols = {col for col, _, _ in formulas}
                inside = [(col, c) for col, c in numerics
                          if c.value != 0 and (col - 1) in fcols and (col + 1) in fcols]
                if inside and len(formulas) >= 0.6 * (len(formulas) + len(inside)):
                    for col, c in inside[:3]:
                        add("warn", "hardcode-in-formula-row", f"{ws.title}!{c.coordinate}",
                            f"constant {c.value!r} sits between formulas "
                            f"{get_column_letter(col - 1)} and {get_column_letter(col + 1)}")
            # numbers stored as text next to numeric data
            if numerics or formulas:
                anchor = [col for col, _ in numerics] + [col for col, _, _ in formulas]
                for col, c in cells:
                    v = c.value
                    if isinstance(v, str) and not v.startswith("=") and NUM_TEXT.match(v) \
                            and any(abs(col - a) <= 2 for a in anchor):
                        add("warn", "number-as-text", f"{ws.title}!{c.coordinate}",
                            f"{v!r} is text; check decimal separator / import")
    wbv = None if quick else open_wb(path, values=True)
    for rule in (QUICK_RULES if quick else FULL_RULES):
        try:
            rule(wb, wbv, add)
        except Exception as e:  # noqa: BLE001 - one broken rule must not sink the report
            add("info", "rule-failed", rule.__name__, f"{type(e).__name__}: {e}")
    if not quick:
        for ws in wbv.worksheets:
            for c in iter_cells(ws):
                v = c.value
                if isinstance(v, str) and v in ERR_TOKENS:
                    add("error", "cached-error", f"{ws.title}!{c.coordinate}",
                        f"{v} (value cached at last Excel save; run calc for a live recalculation)")

    findings.sort(key=lambda f: (LEVEL_ORDER[f["level"]], f["rule"], f["where"]))
    summary = {"errors": sum(1 for f in findings if f["level"] == "error"),
               "warnings": sum(1 for f in findings if f["level"] == "warn"),
               "infos": sum(1 for f in findings if f["level"] == "info"),
               "by_rule": counts, "quick": quick}
    if as_dict:
        return {"file": str(path), "summary": summary, "findings": findings}
    out = [f"LINT {path}  errors={summary['errors']} warnings={summary['warnings']} infos={summary['infos']}"
           + ("  (quick: cached errors not scanned)" if quick else "")]
    for f in findings:
        out.append(f"  {f['level']:<5} {f['rule']:<24} {f['where']:<34} {f['msg']}")
    capped = [f"{r} ({n})" for r, n in counts.items() if n > max_per_rule]
    if capped:
        out.append("  … capped at %d per rule: %s" % (max_per_rule, ", ".join(capped)))
    return "\n".join(out)

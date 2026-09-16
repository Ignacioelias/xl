"""openpyxl workbooks cached by (path, values-mode); invalidated on mtime/size change."""
import time
from pathlib import Path

import openpyxl

_cache = {}


def _key(path, values):
    return (str(Path(path).resolve()).lower(), bool(values))


def open_wb(path, values=False, reload=False, keep_links=True):
    """Open (or fetch from cache) a workbook.

    values=False -> formulas as text (what the model is), values=True -> cached results as
    last saved by Excel. Neither recalculates; use calc() for that.
    """
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(str(p))
    st = p.stat()
    k = _key(p, values)
    ent = _cache.get(k)
    if ent and not reload and ent["mtime"] == st.st_mtime and ent["size"] == st.st_size:
        ent["hits"] += 1
        return ent["wb"]
    t0 = time.perf_counter()
    wb = openpyxl.load_workbook(
        p, data_only=bool(values), keep_links=keep_links,
        keep_vba=p.suffix.lower() == ".xlsm",
    )
    _cache[k] = {"path": str(p.resolve()), "values": bool(values), "mtime": st.st_mtime,
                 "size": st.st_size, "wb": wb, "loaded_at": time.time(),
                 "load_secs": round(time.perf_counter() - t0, 2), "hits": 0}
    return wb


def cache_info():
    return [{k: v for k, v in e.items() if k != "wb"} for e in _cache.values()]


def drop(path=None):
    if path is None:
        n = len(_cache)
        _cache.clear()
        return n
    n = 0
    for values in (False, True):
        if _cache.pop(_key(path, values), None) is not None:
            n += 1
    return n


def resolve_wb(target, values=False):
    """Accept a path or an already-open Workbook."""
    if isinstance(target, openpyxl.Workbook):
        return target
    return open_wb(target, values=values)


def split_ref(ref):
    """'Sheet Name'!A1:B2 or Sheet!A1 or A1:B2 -> (sheet_or_None, range)."""
    ref = ref.strip()
    if "!" in ref:
        sheet, rng = ref.rsplit("!", 1)
        sheet = sheet.strip()
        if sheet.startswith("'") and sheet.endswith("'"):
            sheet = sheet[1:-1].replace("''", "'")
        return sheet, rng.replace("$", "")
    return None, ref.replace("$", "")


def get_sheet(wb, name=None):
    if name is None:
        return wb.active
    if name in wb.sheetnames:
        return wb[name]
    low = {s.lower(): s for s in wb.sheetnames}
    if name.lower() in low:
        return wb[low[name.lower()]]
    cands = [s for s in wb.sheetnames if name.lower() in s.lower()]
    if len(cands) == 1:
        return wb[cands[0]]
    raise KeyError(f"sheet {name!r} not found; sheets: {wb.sheetnames}")

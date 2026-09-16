"""Real Excel engine through a dedicated hidden COM instance (never the user's session).

Discipline from xlsx-toolkit: DispatchEx (private instance), Calculation=xlManual set only
after a workbook exists, CalculateBeforeSave=False, UpdateLinks=0, errors read as ints or
.Text, Chart.Export only after activating sheet and chart and always assert file size.
"""
import ctypes
import json
import re as _re
import os
import subprocess
import sys
import threading
import time
from pathlib import Path

from .paths import TOOLS_DIR, WORK
from .procs import pid_alive, processes_named

XL_MANUAL = -4135
XL_FORMULAS = -4123
XL_CONSTANTS = 2
XL_ERRORS = 16

_app = None
_pid = None


def _pid_of_hwnd(hwnd):
    pid = ctypes.c_ulong()
    ctypes.windll.user32.GetWindowThreadProcessId(int(hwnd), ctypes.byref(pid))
    return pid.value


def _alive(pid):
    return pid_alive(pid)


def windows_of_pid(pid):
    """Top-level windows (class, title, visible) owned by a process: tells you WHICH dialog
    is blocking a hidden Excel instead of guessing."""
    user32 = ctypes.windll.user32
    out = []
    WNDENUMPROC = ctypes.WINFUNCTYPE(ctypes.c_bool, ctypes.c_void_p, ctypes.c_void_p)

    def cb(hwnd, _lp):
        wpid = ctypes.c_ulong()
        user32.GetWindowThreadProcessId(hwnd, ctypes.byref(wpid))
        if wpid.value != pid:
            return True
        cls = ctypes.create_unicode_buffer(128)
        user32.GetClassNameW(hwnd, cls, 128)
        n = user32.GetWindowTextLengthW(hwnd)
        buf = ctypes.create_unicode_buffer(n + 1)
        user32.GetWindowTextW(hwnd, buf, n + 1)
        out.append({"hwnd": int(hwnd), "class": cls.value, "title": buf.value,
                    "visible": bool(user32.IsWindowVisible(hwnd))})
        return True

    user32.EnumWindows(WNDENUMPROC(cb), 0)
    return out


WM_CLOSE = 0x0010
KEEP_TITLES = ("Microsoft Excel",)  # repair / alert prompts: never auto-dismissed, always reported
_guards = {}
_guard_log = {}


def _foreign_dialogs(pid):
    return [w for w in windows_of_pid(pid)
            if w["class"] == "#32770" and w["title"] and w["title"] not in KEEP_TITLES]


def start_dialog_guard(pid, interval=0.7):
    """Add-ins (Arixcel Explorer, seen 2026-09-07) throw a modal on a fresh instance and block
    every COM call forever. This thread closes such dialogs on OUR hidden Excel (by PID) and
    logs what it closed. Repair prompts titled 'Microsoft Excel' are left alone on purpose."""
    if not pid or pid in _guards:
        return
    stop = threading.Event()
    _guards[pid] = stop
    _guard_log.setdefault(pid, [])

    def run():
        user32 = ctypes.windll.user32
        while not stop.is_set() and _alive(pid):
            try:
                for w in _foreign_dialogs(pid):
                    user32.PostMessageW(w["hwnd"], WM_CLOSE, 0, 0)
                    _guard_log[pid].append({"t": time.time(), "title": w["title"]})
            except Exception:
                pass
            stop.wait(interval)
        _guards.pop(pid, None)

    threading.Thread(target=run, name=f"xl-dialog-guard-{pid}", daemon=True).start()


def stop_dialog_guard(pid):
    ev = _guards.get(pid)
    if ev:
        ev.set()


def guard_log(pid):
    return [e["title"] for e in _guard_log.get(pid, [])]


def kill_pid(pid, wait=5.0):
    """Terminate one of OUR hidden Excel processes (never the user's session) and confirm."""
    if not pid or not _alive(pid):
        return False
    subprocess.run(["taskkill", "/PID", str(pid), "/F"], capture_output=True)
    t0 = time.time()
    while time.time() - t0 < wait and _alive(pid):
        time.sleep(0.1)
    return not _alive(pid)


def app():
    global _app, _pid
    if _app is not None:
        try:
            _ = _app.Visible
            return _app
        except Exception:
            _app, _pid = None, None
    import win32com.client as w32

    a = w32.DispatchEx("Excel.Application")
    _pid = _pid_of_hwnd(a.Hwnd)
    a.Visible = False
    a.DisplayAlerts = False
    a.ScreenUpdating = False
    a.EnableEvents = False
    try:
        a.AskToUpdateLinks = False
    except Exception:
        pass
    try:
        # msoAutomationSecurityForceDisable: no Workbook_Open macros, so no VBA MsgBox captioned
        # "Microsoft Excel" (which the dialog guard must leave alone) can wedge the hidden instance
        a.AutomationSecurity = 3
    except Exception:
        pass
    a.Workbooks.Add()  # Calculation cannot be set until a workbook is open
    a.Calculation = XL_MANUAL
    a.CalculateBeforeSave = False
    _app = a
    start_dialog_guard(_pid)
    return a


def excel_pid():
    return _pid if _app is not None else None


def quit_app():
    global _app, _pid
    if _app is None:
        return
    pid = _pid
    try:
        for wb in list(_app.Workbooks):
            wb.Close(SaveChanges=False)
        _app.Quit()
    except Exception:
        pass
    _app, _pid = None, None
    time.sleep(1.0)
    kill_pid(pid)
    stop_dialog_guard(pid)


def _ws(wb, name):
    """Worksheet by name, tolerant of trailing spaces and case (sheet names like 'Inputs & Drivers ')."""
    if name is None:
        return wb.Worksheets(1)
    try:
        return wb.Worksheets(name)
    except Exception:
        pass
    want = name.strip().lower()
    for ws in wb.Worksheets:
        if ws.Name.strip().lower() == want:
            return ws
    cands = [ws for ws in wb.Worksheets if want in ws.Name.strip().lower()]
    if len(cands) == 1:
        return cands[0]
    raise KeyError(f"sheet {name!r} not found; sheets: {[ws.Name for ws in wb.Worksheets]}")


def _open(path, read_only=True):
    p = str(Path(path).resolve())
    if not os.path.exists(p):
        raise FileNotFoundError(p)
    return app().Workbooks.Open(p, UpdateLinks=0, ReadOnly=read_only,
                                IgnoreReadOnlyRecommended=True, Notify=False)


def _formula_snapshot(wb):
    """{sheet: [(area_address, values)]} for formula cells only (ghost used ranges stay cheap)."""
    snap = {}
    for ws in wb.Worksheets:
        try:
            rng = ws.UsedRange.SpecialCells(XL_FORMULAS)
        except Exception:
            continue
        areas = []
        for ar in rng.Areas:
            areas.append((str(ar.Address).replace("$", ""), ar.Value2))
        snap[ws.Name] = areas
    return snap


VOLATILE_RX = _re.compile(r"(?<![A-Z_.])(NOW|TODAY|RAND|RANDBETWEEN|RANDARRAY)\s*\(", _re.I)

# Value2 of an error cell arrives as a VT_ERROR int; map the ones Excel documents, .Text for the rest
ERR_CODES = {-2146826281: "#DIV/0!", -2146826246: "#N/A", -2146826259: "#NAME?", -2146826288: "#NULL!",
             -2146826252: "#NUM!", -2146826265: "#REF!", -2146826273: "#VALUE!"}


def _area_origin(addr):
    """(col0, row0) of an area address as Excel reports it, with whole-column ($A:$A) and
    whole-row ($5:$5) forms handled; the old regex returned None on those and verify crashed."""
    from openpyxl.utils.cell import range_boundaries
    c1, r1, _, _ = range_boundaries(addr.replace("$", "").split(",")[0])
    return (c1 or 1), (r1 or 1)


def _diff_snapshots(before, after, cap=200, wb=None):
    """Cells whose cached value differs after the rebuild. Volatile formulas (NOW, TODAY, RAND…)
    change on every recalculation by definition, so they are counted separately and never make a
    file 'stale'; before this filter every model with =TODAY() on the cover failed verify."""
    from openpyxl.utils import get_column_letter
    changed, volatile = [], 0

    def _formula(sheet, cell):
        if wb is None:
            return ""
        try:
            return str(_ws(wb, sheet).Range(cell).Formula)
        except Exception:
            return ""

    for sheet, areas_b in before.items():
        areas_a = after.get(sheet) or []
        for (addr, vb), (_, va) in zip(areas_b, areas_a):
            if vb == va:
                continue
            c0, row0 = _area_origin(addr)
            pairs = [(addr, vb, va)] if not isinstance(vb, tuple) else [
                (f"{get_column_letter(c0 + j)}{row0 + i}", x, y)
                for i, (rb, ra) in enumerate(zip(vb, va)) for j, (x, y) in enumerate(zip(rb, ra)) if x != y]
            for cell, x, y in pairs:
                if VOLATILE_RX.search(_formula(sheet, cell)):
                    volatile += 1
                    continue
                changed.append({"sheet": sheet, "cell": cell, "before": x, "after": y})
                if len(changed) >= cap:
                    return changed, volatile
    return changed, volatile


def _error_cells(wb, cap):
    """Every error cell, read per Area (one Value2 + one Formula round trip per area) instead of
    one COM call per cell; a 636k-formula model had ~800 round trips here."""
    from openpyxl.utils import get_column_letter
    errors = []
    for ws in wb.Worksheets:
        ur = ws.UsedRange
        for kind in (XL_FORMULAS, XL_CONSTANTS):
            try:
                rng = ur.SpecialCells(kind, XL_ERRORS)
            except Exception:
                continue
            for ar in rng.Areas:
                c0, r0 = _area_origin(str(ar.Address))
                vals = ar.Value2
                forms = ar.Formula if kind == XL_FORMULAS else None
                if not isinstance(vals, tuple):
                    vals, forms = ((vals,),), (((forms,),) if forms is not None else None)
                for i, row in enumerate(vals):
                    for j, v in enumerate(row):
                        # every cell of a SpecialCells(..., xlErrors) area is an error cell
                        cell = f"{get_column_letter(c0 + j)}{r0 + i}"
                        text = ERR_CODES.get(v) if isinstance(v, int) and not isinstance(v, bool) else None
                        if text is None:
                            try:
                                text = str(ws.Range(cell).Text)
                            except Exception:
                                text = f"error {v}"
                        f = forms[i][j] if forms is not None else None
                        errors.append({"sheet": ws.Name, "cell": cell, "error": text,
                                       "formula": (str(f)[:200] if f else None)})
                        if len(errors) >= cap:
                            return errors
    return errors


def calc(path, save=False, cap=200, verify=False, force=False):
    """Full rebuild in the engine; return every error cell. save=True writes the recalculated
    values back: only on a work copy unless force=True (a synced original would be
    uploaded by OneDrive seconds later), always after a backup to the work dir, and refused if
    Excel opened the file read-only (open elsewhere). verify=True also reports which cached
    values changed (the file is stale if any non-volatile one did)."""
    from .paths import backup_to_work, guard_write
    if save:
        guard_write(path, force=force)
    a = app()
    t0 = time.perf_counter()
    wb = _open(path, read_only=not save)
    try:
        if save and wb.ReadOnly:
            raise PermissionError("Excel opened the file read-only (open in another session?); not saving")
        before = _formula_snapshot(wb) if verify else None
        a.CalculateFullRebuild()
        changed, volatile = _diff_snapshots(before, _formula_snapshot(wb), cap, wb) if verify else (None, 0)
        errors = _error_cells(wb, cap)
        saved, backup = False, None
        if save:
            backup = backup_to_work(path)
            wb.Save()
            saved = True
        res = {"file": str(path), "sheets": wb.Worksheets.Count, "error_count": len(errors),
               "capped": len(errors) >= cap, "errors": errors, "saved": saved, "backup": backup,
               "secs": round(time.perf_counter() - t0, 1)}
        if verify:
            res["changed_count"] = len(changed)
            res["changed"] = changed
            res["volatile_skipped"] = volatile
        return res
    finally:
        wb.Close(SaveChanges=False)


def scenario(path, inputs, outputs, full=True):
    """What-if without saving: open read-only, read outputs, set inputs, recalc, read outputs
    again, close without saving. inputs: {"Sheet!A1": value}. outputs: ["Sheet!B2", ...]."""
    from .wbcache import split_ref
    a = app()
    wb = _open(path, read_only=True)
    try:
        def rd(ref):
            s, ad = split_ref(ref)
            return _ws(wb, s).Range(ad).Value2
        if full:
            a.CalculateFullRebuild()
        base = {r: rd(r) for r in outputs}
        for ref, val in inputs.items():
            s, ad = split_ref(ref)
            _ws(wb, s).Range(ad).Value2 = val
        a.CalculateFull()
        after = {r: rd(r) for r in outputs}
        rows = []
        for r in outputs:
            b, x = base[r], after[r]
            delta = (x - b) if isinstance(b, (int, float)) and isinstance(x, (int, float)) else None
            rows.append([r, b, x, delta])
        return {"inputs": inputs, "rows": rows}
    finally:
        wb.Close(SaveChanges=False)


def sweep(path, input_ref, values, outputs):
    """One-variable sweep: for each value of input_ref, recalc and read outputs (no save)."""
    from .wbcache import split_ref
    a = app()
    wb = _open(path, read_only=True)
    try:
        s, ad = split_ref(input_ref)
        cell = _ws(wb, s).Range(ad)
        original = cell.Value2
        rows = []
        for v in values:
            cell.Value2 = v
            a.CalculateFull()
            row = [v]
            for r in outputs:
                os_, oa = split_ref(r)
                row.append(_ws(wb, os_).Range(oa).Value2)
            rows.append(row)
        cell.Value2 = original
        return {"input": input_ref, "outputs": outputs, "rows": rows}
    finally:
        wb.Close(SaveChanges=False)


def render_diff(png, baseline, out=None, threshold=24):
    """Highlight changed pixels in red over a desaturated baseline; returns changed fraction."""
    from PIL import Image, ImageChops, ImageOps
    a = Image.open(png).convert("RGB")
    b = Image.open(baseline).convert("RGB")
    note = None
    if a.size != b.size:
        note = f"size differs: new {a.size} vs baseline {b.size}; baseline resized for the diff"
        b = b.resize(a.size)
    diff = ImageChops.difference(a, b).convert("L").point(lambda v: 255 if v > threshold else 0)
    changed = sum(1 for v in diff.getdata() if v)
    grey = ImageOps.grayscale(b).convert("RGB")
    red = Image.new("RGB", a.size, (220, 30, 30))
    outimg = Image.composite(red, grey, diff)
    out = out or str(Path(png).with_name(Path(png).stem + "_diff.png"))
    outimg.save(out)
    total = a.size[0] * a.size[1]
    return {"diff_png": out, "changed_pixels": changed, "changed_pct": round(100 * changed / total, 3), "note": note}


CF_UNICODETEXT = 13
GMEM_MOVEABLE = 0x0002


def _clip_api():
    """user32/kernel32 with 64-bit-safe signatures: ctypes defaults every return to a C int, which
    truncates HANDLEs and HGLOBALs on x64 (GlobalLock on a truncated handle silently fails)."""
    user32, kernel32 = ctypes.windll.user32, ctypes.windll.kernel32
    user32.OpenClipboard.argtypes = [ctypes.c_void_p]
    user32.GetClipboardData.restype = ctypes.c_void_p
    user32.GetClipboardData.argtypes = [ctypes.c_uint]
    user32.SetClipboardData.restype = ctypes.c_void_p
    user32.SetClipboardData.argtypes = [ctypes.c_uint, ctypes.c_void_p]
    kernel32.GlobalAlloc.restype = ctypes.c_void_p
    kernel32.GlobalAlloc.argtypes = [ctypes.c_uint, ctypes.c_size_t]
    kernel32.GlobalLock.restype = ctypes.c_void_p
    kernel32.GlobalLock.argtypes = [ctypes.c_void_p]
    kernel32.GlobalUnlock.argtypes = [ctypes.c_void_p]
    return user32, kernel32


def _clipboard_text(retries=5):
    """Current clipboard text (None if empty or not text). The render route below has to go
    through the shared Windows clipboard, so at least the user's text is put back afterwards."""
    user32, kernel32 = _clip_api()
    for _ in range(retries):  # another app can hold the clipboard open for a few ms
        if user32.OpenClipboard(None):
            break
        time.sleep(0.05)
    else:
        return None
    try:
        if not user32.IsClipboardFormatAvailable(CF_UNICODETEXT):
            return None
        h = user32.GetClipboardData(CF_UNICODETEXT)
        if not h:
            return None
        p = kernel32.GlobalLock(h)
        try:
            return ctypes.wstring_at(p) if p else None
        finally:
            kernel32.GlobalUnlock(h)
    finally:
        user32.CloseClipboard()


def _set_clipboard_text(text, retries=5):
    user32, kernel32 = _clip_api()
    data = text.encode("utf-16-le") + b"\x00\x00"
    h = kernel32.GlobalAlloc(GMEM_MOVEABLE, len(data))
    if not h:
        return False
    p = kernel32.GlobalLock(h)
    ctypes.memmove(p, data, len(data))
    kernel32.GlobalUnlock(h)
    for _ in range(retries):
        if user32.OpenClipboard(None):
            break
        time.sleep(0.05)
    else:
        return False
    try:
        user32.EmptyClipboard()
        return bool(user32.SetClipboardData(CF_UNICODETEXT, h))
    finally:
        user32.CloseClipboard()


def _safe_name(s):
    return _re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", str(s)).strip() or "sheet"


def render(path, ref=None, out=None, max_cells=40000, diff=None):
    """Render a range (default: used range of first sheet) to PNG via CopyPicture + chart export.
    The clipboard is shared with the user's session: whatever text they had is restored, and the
    pasted picture is checked against the range's size so a Ctrl+C between our copy and our
    paste fails loudly instead of exporting the user's clipboard as 'the range'."""
    from .wbcache import split_ref

    a = app()
    wb = _open(path, read_only=True)
    clip_before = None
    try:
        clip_before = _clipboard_text()
    except Exception:
        pass
    try:
        sheet, rng = split_ref(ref) if ref else (None, None)
        ws = _ws(wb, sheet)
        r = ws.Range(rng) if rng else ws.UsedRange
        n = r.Rows.Count * r.Columns.Count
        if n > max_cells:
            raise ValueError(f"{n:,} cells; render a smaller range (max {max_cells:,})")
        out = str(Path(out or (WORK / f"render_{_safe_name(ws.Name)}_{int(time.time())}.png")).resolve())
        ws.Activate()
        last_err = None
        for attempt in range(4):  # the clipboard is shared with the user's Office session
            try:
                a.CutCopyMode = False
                r.CopyPicture(Appearance=1, Format=2)  # xlScreen, xlBitmap
                last_err = None
                break
            except Exception as e:  # noqa: BLE001
                last_err = e
                time.sleep(0.6 * (attempt + 1))
        if last_err is not None:
            # clipboard route dead (busy clipboard, remote session): print the range to PDF and rasterise
            try:
                import fitz  # PyMuPDF
            except ImportError:
                raise RuntimeError(f"CopyPicture failed 4 times (clipboard busy?) and PyMuPDF is not installed: {last_err}")
            pdf = str(Path(out).with_suffix(".pdf"))
            ws.PageSetup.PrintArea = str(r.Address)
            ws.PageSetup.Zoom = False
            ws.PageSetup.FitToPagesWide = 1
            ws.PageSetup.FitToPagesTall = 1
            wb.ExportAsFixedFormat(0, pdf, 0, True, False, 1, 1, False)  # xlTypePDF, quality standard
            doc = fitz.open(pdf)
            pix = doc[0].get_pixmap(matrix=fitz.Matrix(2, 2))
            pix.save(out)
            doc.close()
            try:
                os.unlink(pdf)
            except OSError:
                pass
            size = os.path.getsize(out)
            res = {"png": out, "bytes": size, "range": f"{ws.Name}!{str(r.Address).replace(chr(36), str())}", "method": "pdf-fallback"}
            if diff:
                res.update(render_diff(out, diff))
            return res
        co = ws.ChartObjects().Add(r.Left, r.Top, max(r.Width, 10), max(r.Height, 10))
        try:
            co.Activate()
            ch = co.Chart
            try:
                ch.ChartArea.Format.Line.Visible = 0
            except Exception:
                pass
            ch.Paste()
            # the paste must be OUR picture: compare its size with the range (15% tolerance)
            try:
                pic = ch.Shapes(ch.Shapes.Count) if ch.Shapes.Count else None
                pw, ph = (float(pic.Width), float(pic.Height)) if pic is not None else (None, None)
            except Exception:
                pw = ph = None
            if pw is not None and (abs(pw - r.Width) > 0.15 * r.Width + 4 or abs(ph - r.Height) > 0.15 * r.Height + 4):
                raise RuntimeError(f"clipboard changed between copy and paste (pasted {pw:.0f}x{ph:.0f} pt, "
                                   f"range {r.Width:.0f}x{r.Height:.0f} pt); nothing exported, retry")
            ch.Export(out, "PNG")
        finally:
            co.Delete()
        size = os.path.getsize(out) if os.path.exists(out) else 0
        method = "chart-export"
        if size == 0:
            from PIL import ImageGrab
            r.CopyPicture(Appearance=1, Format=2)
            im = ImageGrab.grabclipboard()
            if im is None:
                raise RuntimeError("render wrote 0 bytes and the clipboard held no image")
            im.save(out)
            size = os.path.getsize(out)
            method = "clipboard"
        res = {"png": out, "bytes": size, "range": f"{ws.Name}!{str(r.Address).replace(chr(36), str())}", "method": method}
        if diff:
            res.update(render_diff(out, diff))
        return res
    finally:
        wb.Close(SaveChanges=False)
        try:
            a.CutCopyMode = False
            if clip_before:
                _set_clipboard_text(clip_before)
        except Exception:
            pass


def check_open(path, timeout=90):
    """Law 7: open in a throwaway Excel with DisplayAlerts ON. A repair prompt blocks, the
    timeout fires, the blocking dialog is named, and the file is reported as NOT verified.
    The throwaway Excel is always terminated by PID afterwards (Quit alone leaves it alive).
    If the child never wrote its pidfile (an add-in modal during start-up), the instance is
    found by elimination: a hidden EXCEL.EXE created after we started that is not the kernel's."""
    pidfile = WORK / f"check_{os.getpid()}_{int(time.time() * 1000)}.pid"
    t_spawn = time.time()
    excel_before = set(processes_named("EXCEL.EXE"))
    proc = subprocess.Popen(
        [sys.executable, "-m", "xl.excelcom", "--check", str(Path(path).resolve()), str(pidfile)],
        cwd=str(TOOLS_DIR), stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True,
        creationflags=subprocess.CREATE_NO_WINDOW,
    )

    def child_excel_pid():
        try:
            if pidfile.exists():
                return int(pidfile.read_text().strip())
        except ValueError:
            pass
        # fallback: a new hidden EXCEL.EXE, started after us, that is not the kernel's own
        for pid, t in processes_named("EXCEL.EXE").items():
            if pid in excel_before or pid == _pid or t is None or t < t_spawn - 2:
                continue
            if any(w["visible"] and w["class"] == "XLMAIN" for w in windows_of_pid(pid)):
                continue  # a visible Excel is the user's session, never ours
            return pid
        return None

    t0 = time.time()
    while time.time() - t0 < 20 and child_excel_pid() is None and proc.poll() is None:
        time.sleep(0.2)
    start_dialog_guard(child_excel_pid())
    try:
        out, _ = proc.communicate(timeout=max(5, timeout - (time.time() - t0)))
        last = [ln for ln in out.strip().splitlines() if ln.startswith("{")]
        res = json.loads(last[-1]) if last else {"opened": False, "blocked": False, "msg": out.strip()[-800:]}
    except subprocess.TimeoutExpired:
        pid = child_excel_pid()
        wins = [w for w in windows_of_pid(pid) if w["title"]] if pid else []
        dialogs = [w for w in wins if w["class"] in ("#32770", "bosa_sdm_XL9", "NUIDialog")] or wins
        proc.kill()
        res = {"opened": False, "blocked": True, "excel_pid": pid,
               "dialogs": dialogs[:6],
               "msg": (f"Excel did not finish opening within {timeout}s; blocking window(s): "
                       + "; ".join(f"[{w['class']}] {w['title']}" for w in dialogs[:4])
                       if dialogs else
                       f"Excel did not finish opening within {timeout}s and showed no titled window "
                       f"(slow open or a hidden modal). NOT verified.")}
    finally:
        pid = child_excel_pid()
        killed = kill_pid(pid) if pid else False
        try:
            pidfile.unlink()
        except OSError:
            pass
    res["excel_terminated"] = killed
    res["dismissed_dialogs"] = guard_log(pid) if pid else []
    stop_dialog_guard(pid)
    return res


def _check_main(path, pidfile):
    import pythoncom
    import win32com.client as w32

    pythoncom.CoInitialize()
    a = w32.DispatchEx("Excel.Application")
    Path(pidfile).write_text(str(_pid_of_hwnd(a.Hwnd)))  # first thing: the parent can only kill what it knows
    a.Visible = False
    a.DisplayAlerts = True
    try:
        a.AutomationSecurity = 3  # the law-7 test is about repair prompts, not about running macros
    except Exception:
        pass
    a.Workbooks.Add()
    a.Calculation = XL_MANUAL
    t0 = time.perf_counter()
    wb = a.Workbooks.Open(path, UpdateLinks=0, ReadOnly=True, IgnoreReadOnlyRecommended=True)
    info = {"opened": True, "blocked": False, "read_only": bool(wb.ReadOnly), "sheets": wb.Worksheets.Count,
            "name": wb.Name, "secs": round(time.perf_counter() - t0, 1)}
    wb.Close(SaveChanges=False)
    a.DisplayAlerts = False  # otherwise Quit() asks to save the scratch workbook and hangs hidden forever
    for w in list(a.Workbooks):
        w.Close(SaveChanges=False)
    a.Quit()
    del a
    pythoncom.CoUninitialize()
    print(json.dumps(info), flush=True)


if __name__ == "__main__":
    if len(sys.argv) >= 4 and sys.argv[1] == "--check":
        _check_main(sys.argv[2], sys.argv[3])

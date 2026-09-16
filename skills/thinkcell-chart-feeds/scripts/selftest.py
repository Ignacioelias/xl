"""tcfeed selftest.

  py -3.14 selftest.py [--dir DIR] [--session NAME]

0. Semantics probe: every builder's formula is evaluated by Excel on edge-case data (text in
   the data, a duplicate and a text year, blanks, missing matches, wildcards, a zero
   denominator) and compared with the Python mirror.
1. Builds a synthetic "model" in a dedicated hidden Excel: a quarterly sheet with a space in
   its name, an annual sheet with formulas and an actual/forecast flag, an inputs sheet, and a
   pre-existing feed tab with two blocks, two hidden sheet-scoped names standing in for
   think-cell links, a shape, and a formula that points below the insertion point. Saved with
   manual calculation and iterate="1".
2. Runs the pipeline on a copy from examples/sample_spec.json (a standard block inserted above
   an existing block, a waterfall appended), with the engine check.
3. Asserts every verification check passed and a set of specific facts.
4. Negative controls: tampered copies that each verification check must reject.
5. CLI paths: verify, compute, the lock guard, spec errors, a failure before and after the save.
Stops the xl session at the end. Output folder: %LOCALAPPDATA%/thinkcell-chart-feeds/selftest
(wiped first; --dir to change). Exit 0 = all pass.
"""
import argparse
import hashlib
import json
import os
import re
import shutil
import sys
import time
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from tcfeed import a1, pipeline, reader, package, xlbridge  # noqa: E402
from tcfeed import verify as V  # noqa: E402
from tcfeed import xlsem as X  # noqa: E402

T1, T2 = "___thinkcellTEST0001", "___thinkcellTEST0002"
RESULTS = []


def ok(name, cond, detail=""):
    RESULTS.append({"test": name, "pass": bool(cond), "detail": str(detail)[:300]})
    print(f"{'PASS' if cond else 'FAIL'}  {name}" + (f"  [{str(detail)[:200]}]" if detail else ""), flush=True)


def sha(path):
    return hashlib.sha256(open(path, "rb").read()).hexdigest()


# ------------------------------------------------------------------ synthetic model
def _row(ws, r, c0, vals):
    ws.Range(ws.Cells(r, c0), ws.Cells(r, c0 + len(vals) - 1)).Formula = (tuple(vals),)


def build_model(path):
    import win32com.client as w32
    xc, pr = xlbridge.excelcom(), xlbridge.procs()
    app = w32.DispatchEx("Excel.Application")
    pid = xc._pid_of_hwnd(app.Hwnd)
    xc.start_dialog_guard(pid)
    app.Visible = False
    app.DisplayAlerts = False
    app.ScreenUpdating = False
    app.EnableEvents = False
    wb = None
    try:
        wb = app.Workbooks.Add()
        while wb.Worksheets.Count < 4:
            wb.Worksheets.Add(After=wb.Worksheets(wb.Worksheets.Count))
        for i, n in enumerate(["Inputs", "Calc Q", "Calc_A", "Feed"], 1):
            wb.Worksheets(i).Name = n
        inp, q, an, fd = (wb.Worksheets(n) for n in ("Inputs", "Calc Q", "Calc_A", "Feed"))

        inp.Range("B2").Value = "Scenario"
        inp.Range("C2").Value = "Base case"
        inp.Range("B3").Value = "Currency"
        inp.Range("C3").Value = "EUR m"

        # quarterly grid E:AB = 2021Q1..2026Q4
        q.Range("B1").Value = "QUARTERLY MODEL (EUR m)"
        q.Range("B2").Value = "Year"
        q.Range("B3").Value = "Quarter"
        _row(q, 2, 5, [2021 + i // 4 for i in range(24)])
        _row(q, 3, 5, [i % 4 + 1 for i in range(24)])
        labels = {5: "Segment A", 6: "Segment B", 7: "Segment C", 8: "Total revenue",
                  10: "Customers, end of quarter ('000)", 11: "Customers, sparse ('000)"}
        for r, t in labels.items():
            q.Cells(r, 2).Value = t
        _row(q, 5, 5, [round(10 + 0.37 * i + (i % 4) * 0.113, 4) for i in range(24)])
        _row(q, 6, 5, [round(5.5 + 0.217 * i, 4) for i in range(24)])
        _row(q, 7, 5, [0.0 if i < 8 else round(1.3 + 0.51 * (i - 8), 4) for i in range(24)])
        _row(q, 8, 5, [f"=SUM({a1.col_letter(5 + i)}5:{a1.col_letter(5 + i)}7)" for i in range(24)])
        _row(q, 10, 5, [100 + 3 * i for i in range(24)])
        _row(q, 11, 5, [None if i in (11, 19) else 50 + 2 * i for i in range(24)])   # blank Q4 of 2023, 2025

        # annual grid E:N = 2021..2030, formulas for 2021-26, values after
        an.Range("B1").Value = "ANNUAL MODEL (EUR m)"
        an.Range("B2").Value = "Year"
        an.Range("B3").Value = "Actual flag"
        _row(an, 2, 5, list(range(2021, 2031)))
        _row(an, 3, 5, [1 if y <= 2023 else 0 for y in range(2021, 2031)])
        for r, lab in ((5, "Segment A"), (6, "Segment B"), (7, "Segment C"), (9, "Total revenue")):
            an.Cells(r, 2).Value = lab
            an.Cells(r, 3).Value = "Total" if r == 9 else "Revenue"
        for r, base in ((5, 70.0), (6, 30.0), (7, 25.0)):
            vals = [f"=SUMIFS('Calc Q'!$E${r}:$AB${r},'Calc Q'!$E$2:$AB$2,{a1.col_letter(5 + k)}$2)"
                    for k in range(6)]
            vals += [base + 2.5 * k for k in range(4)]
            _row(an, r, 5, vals)
        _row(an, 9, 5, [f"=SUM({a1.col_letter(5 + k)}5:{a1.col_letter(5 + k)}7)" for k in range(10)])

        # the feed tab as it exists before the pipeline runs
        fd.Range("B1").Value = "CHART FEEDS | synthetic test model"
        fd.Range("B1").Font.Bold = True
        fd.Range("B3").Value = "Existing block 1 | total revenue (EUR m)"
        fd.Range("B3").Font.Bold = True
        _row(fd, 4, 3, [f"'FY{y % 100:02d}" for y in range(2021, 2027)])
        fd.Range("C4:H4").Font.Bold = True
        fd.Range("C4:H4").HorizontalAlignment = -4152
        fd.Range("B6").Value = "Total revenue"
        _row(fd, 6, 3, [f"=Calc_A!{a1.col_letter(5 + k)}9" for k in range(6)])
        fd.Range("B7").Value = "Segment A"
        _row(fd, 7, 3, [f"=Calc_A!{a1.col_letter(5 + k)}5" for k in range(6)])
        fd.Range("C6:H7").NumberFormat = "#,##0.0;(#,##0.0);0.0"
        fd.Range("B8").Value = "Check (=0) - total less the three segments"
        _row(fd, 8, 3, [f"={a1.col_letter(3 + k)}6-Calc_A!{a1.col_letter(5 + k)}5"
                        f"-Calc_A!{a1.col_letter(5 + k)}6-Calc_A!{a1.col_letter(5 + k)}7" for k in range(6)])
        fd.Range("B8:H8").Font.Italic = True
        fd.Range("C8:H8").NumberFormat = "#,##0.000;(#,##0.000);0.000"
        fd.Range("B9").Value = "Memo - FY2021 customers, read from block 2 below"
        fd.Range("C9").Formula = "=C14"
        fd.Range("C9").NumberFormat = "#,##0;(#,##0);0"
        fd.Range("B11").Value = "Existing block 2 | customers at year end ('000)"
        fd.Range("B11").Font.Bold = True
        _row(fd, 12, 3, [f"'FY{y % 100:02d}" for y in range(2021, 2027)])
        fd.Range("C12:H12").Font.Bold = True
        fd.Range("B14").Value = "Customers at year end"
        _row(fd, 14, 3, [f"=LOOKUP(2,1/('Calc Q'!$E$2:$AB$2={y}),'Calc Q'!$E$10:$AB$10)" for y in range(2021, 2027)])
        fd.Range("B15").Value = "Memo - customers, twice"
        fd.Range("C15").Formula = "=C14*2"
        fd.Names.Add(Name=T1, RefersTo="=Feed!$B$4:$H$7", Visible=False)
        fd.Names.Add(Name=T2, RefersTo="=Feed!$B$12:$H$14", Visible=False)
        fd.Shapes.AddShape(1, fd.Range("J12").Left, fd.Range("J12").Top, 60, 18)

        app.CalculateFull()
        app.Iteration = True
        app.Calculation = -4135
        wb.SaveAs(os.path.normpath(path), FileFormat=51)
    finally:
        try:
            if wb is not None:
                wb.Close(SaveChanges=False)
        finally:
            app.Quit()
            del app
            t0 = time.time()
            while time.time() - t0 < 20 and pr.pid_alive(pid):
                time.sleep(0.25)
            if pr.pid_alive(pid) and not any(w["visible"] and w["class"] == "XLMAIN" for w in xc.windows_of_pid(pid)):
                xc.kill_pid(pid)
            xc.stop_dialog_guard(pid)


# ------------------------------------------------------------------ semantics probe
class _ProbeCx:
    """Minimal context: builders read `data` and resolve keys to fixed cells."""

    def __init__(self, data, year, keys):
        from tcfeed import builders as B
        self._B, self.data, self.keys = B, data, keys
        self.periods = {"year": B.Period(year)}

    def period(self, tok):
        return self._B.resolve_period(tok, self.periods, "probe")

    def model(self, sheet, r, c):
        return self.data.get((sheet, r, c))

    def formula_row(self, sheet, r):
        return r

    def key_addr(self, k):
        return self.keys[k][0]

    def key_value(self, k):
        return self.keys[k][1]


def semantics_probe():
    """Excel evaluates each builder on edge-case data; the Python mirror must agree."""
    import win32com.client as w32
    from tcfeed import builders as B
    xc, pr = xlbridge.excelcom(), xlbridge.procs()
    rows = {
        1: [2024, 2025, "FY25", 2025],            # duplicate year, text year
        2: [1.5, 2.5, 3.5, 4.5],
        3: [1, "abc", 2, 3],                       # text in the data
        4: [10, 20, 30, None],                     # blank at the year's last column
        5: [5, 6, 7, 8], 6: [50, 60, 70, 80], 7: [500, 600, 700, 800],
        8: [1, 0, None, 0],                        # actual flag
    }
    labels = {5: "Revenue", 6: "revenue other", 7: "Cost"}
    data = {}
    for r, vals in rows.items():
        for k, v in enumerate(vals):
            if v is not None:
                data[("P", r, 3 + k)] = float(v) if isinstance(v, int) else v
    for r, t in labels.items():
        data[("P", r, 2)] = t
    data[("P", 1, 11)] = "some text"
    grid = B.Grid("g", "P", 1, 3, 6)
    ctx = {"grids": {"g": grid}}
    kv = {"a": ("C2", 1.5), "blank": ("L1", None), "txt": ("D3", "abc"), "d2": ("D2", 2.5), "e2": ("E2", 3.5)}
    cases = [
        ("sumifs: duplicate year summed", {"fn": "sumifs", "grid": "g", "row": 2}, 2025),
        ("sumifs: text in data skipped", {"fn": "sumifs", "grid": "g", "row": 3}, 2025),
        ("sumifs: scale and sign", {"fn": "sumifs", "grid": "g", "rows": [2, 5], "scale": 0.001, "sign": -1}, 2025),
        ("window: text in data is #VALUE!", {"fn": "window", "grid": "g", "row": 3, "y0": 2024}, 2025),
        ("window: text year excluded", {"fn": "window", "grid": "g", "rows": [2, 4], "y0": 2024}, 2025),
        ("window: no lower bound", {"fn": "window", "grid": "g", "row": 2}, 2024),
        ("colcount: text year not counted", {"fn": "colcount", "grid": "g", "y0": 2024, "y1": 2026, "expect": 0}, 2025),
        ("index_match: first column of the year", {"fn": "index_match", "grid": "g", "row": 2}, 2025),
        ("index_match: no match is #N/A", {"fn": "index_match", "grid": "g", "row": 2}, 2030),
        ("index_match: offset past the row is #REF!", {"fn": "index_match", "grid": "g", "row": 2, "offset": 5}, 2025),
        ("index_match: text passes through", {"fn": "index_match", "grid": "g", "row": 3}, 2025),
        ("index_match: text scaled is #VALUE!", {"fn": "index_match", "grid": "g", "row": 3, "scale": 1000}, 2025),
        ("year_end: last column of the year", {"fn": "year_end", "grid": "g", "row": 2}, 2025),
        ("year_end: blank cell reads 0", {"fn": "year_end", "grid": "g", "row": 4}, 2025),
        ("year_end: no match is #N/A", {"fn": "year_end", "grid": "g", "row": 2}, 2031),
        ("sumifs_label: wildcard, case-insensitive", {"fn": "sumifs_label", "grid": "g", "label_col": "B",
                                                      "label": "REV*", "first_row": 5, "last_row": 7}, 2025),
        ("sumifs_label: exact label", {"fn": "sumifs_label", "grid": "g", "label_col": "B",
                                       "label": "revenue", "first_row": 5, "last_row": 7}, 2025),
        ("link: blank reads 0", {"fn": "link", "ref": "P!$L$1"}, 2025),
        ("link: text passes through", {"fn": "link", "ref": "P!$K$1"}, 2025),
        ("ratio: zero denominator is #DIV/0!", {"fn": "ratio", "num": "a", "den": "blank"}, 2025),
        ("ratio: blank_if_zero gives empty text", {"fn": "ratio", "num": "a", "den": "blank", "blank_if_zero": True}, 2025),
        ("combine: text operand is #VALUE!", {"fn": "combine", "terms": [{"key": "a"}, {"key": "txt", "coef": -1}]}, 2025),
        ("combine: single key passes text", {"fn": "combine", "terms": [{"key": "txt"}]}, 2025),
        ("combine: coefficient and builder term", {"fn": "combine", "terms": [
            {"key": "a", "coef": 0.5}, {"fn": "index_match", "grid": "g", "row": 2, "coef": -2}]}, 2025),
        ("tie: total less SUM(parts)", {"fn": "tie", "total": "a", "parts": ["d2", "e2", "txt"]}, 2025),
        ("fy_label: forecast flag", {"fn": "fy_label", "actual_flag": {"grid": "g", "row": 8}}, 2025),
        ("fy_label: actual flag", {"fn": "fy_label", "actual_flag": {"grid": "g", "row": 8}}, 2024),
        ("fy_label: no match is #N/A", {"fn": "fy_label", "actual_flag": {"grid": "g", "row": 8}}, 2027),
    ]
    app = w32.DispatchEx("Excel.Application")
    pid = xc._pid_of_hwnd(app.Hwnd)
    xc.start_dialog_guard(pid)
    app.Visible = False
    app.DisplayAlerts = False
    wb = None
    out = []
    try:
        wb = app.Workbooks.Add()
        ws = wb.Worksheets(1)
        ws.Name = "P"
        for r, vals in rows.items():
            _row(ws, r, 3, [("'" + v) if isinstance(v, str) else v for v in vals])
        for r, t in labels.items():
            ws.Cells(r, 2).Value = t
        ws.Range("K1").Value = "some text"
        built = []
        for i, (name, spec, year) in enumerate(cases):
            b = B.make(spec, ctx, name)
            cx = _ProbeCx(data, year, kv)
            ws.Cells(20 + i, 8).Formula = "=" + b.formula(cx)
            built.append((name, b, cx))
        app.Calculate()
        for i, (name, b, cx) in enumerate(built):
            got = ws.Cells(20 + i, 8).Value2
            if isinstance(got, int) and not isinstance(got, bool) and got in xc.ERR_CODES:
                got = X.XlError(xc.ERR_CODES[got])
            want = b.value(cx)
            out.append((name, X.same_value(got, want, rel=1e-12, abs_tol=1e-12), f"excel={got!r} mirror={want!r}"))
    finally:
        try:
            if wb is not None:
                wb.Close(SaveChanges=False)
        finally:
            app.Quit()
            del app
            t0 = time.time()
            while time.time() - t0 < 20 and pr.pid_alive(pid):
                time.sleep(0.25)
            if pr.pid_alive(pid) and not any(w["visible"] and w["class"] == "XLMAIN" for w in xc.windows_of_pid(pid)):
                xc.kill_pid(pid)
            xc.stop_dialog_guard(pid)
    return out


# ------------------------------------------------------------------ tampering helpers
def zip_patch(src, dst, part, fn):
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == part:
                data = fn(data.decode("utf-8")).encode("utf-8")
            zout.writestr(item, data)


def openpyxl_edit(src, dst, fn):
    import openpyxl
    wb = openpyxl.load_workbook(src)
    fn(wb["Feed"])
    wb.save(dst)


# ------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser()
    # outside the skill folder (the folder is wiped first, so never point this at the skill)
    ap.add_argument("--dir", default=os.path.join(os.environ.get("LOCALAPPDATA") or os.path.expanduser("~"),
                                                  "thinkcell-chart-feeds", "selftest"))
    ap.add_argument("--session", default="tcpipe")
    args = ap.parse_args()
    test, session = args.dir, args.session
    t_start = time.time()
    if os.path.exists(test):
        shutil.rmtree(test)
    os.makedirs(test)
    cached = xlbridge.cached()
    try:
        # ---- 0. builder semantics, Excel against the mirror
        for name, good, detail in semantics_probe():
            ok(f"semantics: {name}", good, detail)

        # ---- 1. synthetic model
        src = os.path.join(test, "synthetic_model_source.xlsx")
        build_model(src)
        model = os.path.join(test, "model_under_test.xlsx")
        shutil.copyfile(src, model)
        spec = os.path.join(test, "selftest_spec.json")
        shutil.copyfile(os.path.join(HERE, "examples", "sample_spec.json"), spec)
        calc_src = cached.read_calcpr(src)
        ok("synthetic model saved with manual calculation and iterate=1",
           calc_src.get("calcMode") == "manual" and calc_src.get("iterate") == "1", calc_src)
        names_src = package.local_names(src, "Feed")
        ok("feed tab carries two hidden sheet-scoped names", set(names_src) == {T1, T2}
           and all(v["hidden"] for v in names_src.values()), names_src)

        # ---- 2. dry run and readers
        code = pipeline.main(["plan", spec, "--session", session])
        ok("plan (dry run) exits 0", code == 0, f"exit {code}")
        ok("dry run left the workbook byte-identical", sha(model) == sha(src))
        _, plan = pipeline.load_plan(spec)
        needs = plan.needs()
        ms = reader.fetch_stream(model, needs)
        mx = reader.fetch_xl(model, needs, session)
        ok("stream reader and xl-kernel reader return the same cached values",
           ms.as_dict() == mx.as_dict() and len(ms) > 0, f"{len(ms)} cells")

        # ---- 3. the run
        run_dir = os.path.join(test, "run")
        code = pipeline.main(["run", spec, "--engine-check", "--run-dir", run_dir, "--session", session])
        ok("run exits 0: built, injected, calcPr restored, verified", code == 0, f"exit {code}")
        rep = json.load(open(os.path.join(run_dir, "report.json"), encoding="utf-8"))
        status = {i["check"]: i["status"] for i in rep.get("verify", [])}
        for k in ("layout", "injected", "checks_zero", "names", "untouched_feed", "untouched_model",
                  "errors", "merged", "calcpr", "package", "check_open", "engine"):
            ok(f"verify.{k} = PASS", status.get(k) == "PASS", status.get(k))
        n = rep["build"]["insertions"][0]["rows"]
        names_run = package.local_names(model, "Feed")
        ok("think-cell name above the insertion unchanged",
           a1.norm_ref_text(names_run[T1]["text"]) == a1.norm_ref_text(names_src[T1]["text"]),
           names_run[T1]["text"])
        ok(f"think-cell name below the insertion re-pointed by exactly {n} rows",
           a1.norm_ref_text(names_run[T2]["text"]) == a1.norm_ref_text(f"Feed!B{12 + n}:H{14 + n}"),
           names_run[T2]["text"])
        ok("calcPr identical to the source after the run (iterate=1 kept)",
           cached.read_calcpr(model) == calc_src, cached.read_calcpr(model))
        vals, forms = reader.scan_sheet(model, "Feed")
        ok("existing formula pointing below the insertion was re-pointed by Excel",
           forms.get((9, 3), "").replace("$", "") == f"=C{14 + n}", forms.get((9, 3)))
        pj = json.load(open(os.path.join(run_dir, "plan.json"), encoding="utf-8"))
        bridge = next(b for b in pj["blocks"] if b["id"] == "bridge")
        segs = [r["row"] for r in bridge["rows"] if r["kind"] == "segment"]
        ok("waterfall closing column carries the literal 'e' on the last segment row only",
           vals.get((segs[-1], bridge["last_col"])) == "e"
           and all((r, bridge["last_col"]) not in vals for r in segs[:-1]))
        inputs_row = next(r["row"] for r in bridge["rows"] if r["kind"] == "inputs")
        launch = forms.get((segs[-1], bridge["first_col"] + 3), "")
        ok("waterfall step '$open+1' reads the opening-year input cell",
           f"$C${inputs_row}+1" in launch, launch[:120])
        hdr = next(r["row"] for r in bridge["rows"] if r["kind"] == "header")
        ok("waterfall headers follow the input cells", vals.get((hdr, 3)) == "FY22"
           and vals.get((hdr, bridge["last_col"])) == "FY26", (vals.get((hdr, 3)), vals.get((hdr, bridge["last_col"]))))
        rev = next(b for b in pj["blocks"] if b["id"] == "revenue")
        rhdr = next(r["row"] for r in rev["rows"] if r["kind"] == "header")
        ok("standard block headers carry the model's A/F flag",
           [vals.get((rhdr, c)) for c in range(3, 9)] == ["FY21A", "FY22A", "FY23A", "FY24F", "FY25F", "FY26F"],
           [vals.get((rhdr, c)) for c in range(3, 9)])
        sparse = next(r["row"] for r in rev["rows"] if r["key"] == "sparse_ye")
        ok("a blank Q4 reads 0 through LOOKUP (FY2023 sparse stock)", vals.get((sparse, 5)) == 0,
           vals.get((sparse, 5)))
        ev = rep["build"]["excel_cached_on_entry"]
        print(f"INFO  Excel's own cached value for new formulas before injection: {ev}")

        # ---- 4. negative controls: each tampered copy must fail its check
        neg = os.path.join(test, "negative")
        os.makedirs(neg)
        info = json.load(open(os.path.join(run_dir, "run.json"), encoding="utf-8"))
        _, vplan = pipeline.load_plan(info["spec"], source=info["backup"])
        values, errors, problems = pipeline.do_compute(vplan, info["backup"], "stream", session, print)

        def verify(path):
            res = V.run(vplan, info["backup"], path, values, errors, info["calcpr_expected"],
                        session=session, open_check=False, log=print)
            return {i["check"]: i["status"] for i in res.items}

        base = verify(model)
        ok("control: the delivered file passes verification without check-open",
           all(s in ("PASS", "SKIPPED") for s in base.values()), base)

        first_seg = next(r["row"] for r in rev["rows"] if r["key"] == "seg_a")
        addr = a1.addr(4, first_seg)
        t1 = os.path.join(neg, "wrong_cached_value.xlsx")
        cached.inject(model, "Feed", {addr: values[(first_seg, 4)] + 1.0}, out=t1)
        s = verify(t1)
        ok("negative: a wrong cached value fails 'injected'", s["injected"] == "FAIL", s["injected"])

        chk_row = next(r["row"] for r in rev["rows"] if r["kind"] == "check")
        t2 = os.path.join(neg, "check_not_zero.xlsx")
        cached.inject(model, "Feed", {a1.addr(3, chk_row): 0.5}, out=t2)
        s = verify(t2)
        ok("negative: a check cell reading 0.5 fails 'checks_zero'", s["checks_zero"] == "FAIL", s["checks_zero"])

        t3 = os.path.join(neg, "calcpr_full_calc_on_load.xlsx")
        cached.calcpr(model, set={"fullCalcOnLoad": "1"}, out=t3)
        s = verify(t3)
        ok("negative: fullCalcOnLoad added fails 'calcpr'", s["calcpr"] == "FAIL", s["calcpr"])

        spacer = next(r["row"] for r in rev["rows"] if r["kind"] == "spacer")
        t4 = os.path.join(neg, "spacer_written.xlsx")
        openpyxl_edit(model, t4, lambda ws: ws.cell(spacer, 5, "x"))
        s = verify(t4)
        ok("negative: text in the spacer row fails 'layout'", s["layout"] == "FAIL", s["layout"])

        seg_b = next(r["row"] for r in rev["rows"] if r["key"] == "seg_b")

        def swap(ws):
            a, b = ws.cell(first_seg, 2).value, ws.cell(seg_b, 2).value
            ws.cell(first_seg, 2).value, ws.cell(seg_b, 2).value = b, a
        t5 = os.path.join(neg, "labels_drifted.xlsx")
        openpyxl_edit(model, t5, swap)
        s = verify(t5)
        ok("negative: labels out of step with their formulas fail 'layout'", s["layout"] == "FAIL", s["layout"])

        t6 = os.path.join(neg, "name_broken.xlsx")
        zip_patch(model, t6, "xl/workbook.xml",
                  lambda x: re.sub(r"(<definedName[^>]*name=\"%s\"[^>]*>)[^<]*" % T2, r"\1Feed!#REF!", x))
        s = verify(t6)
        ok("negative: a broken think-cell name fails 'names'", s["names"] == "FAIL", s["names"])

        t7 = os.path.join(neg, "existing_block_edited.xlsx")
        cached.inject(model, "Feed", {"C6": 999.0}, out=t7)
        s = verify(t7)
        ok("negative: a changed value above the new blocks fails 'untouched_feed'",
           s["untouched_feed"] == "FAIL", s["untouched_feed"])

        t8 = os.path.join(neg, "model_value_changed.xlsx")
        cached.inject(model, "Calc_A", {"E9": 1.0}, out=t8)
        s = verify(t8)
        ok("negative: a changed cached value on a model sheet fails 'untouched_model'",
           s["untouched_model"] == "FAIL", s["untouched_model"])

        t9 = os.path.join(neg, "trash_part.xlsx")
        shutil.copyfile(model, t9)
        with zipfile.ZipFile(t9, "a") as z:
            z.writestr("[trash]/0000.dat", b"x")
        s = verify(t9)
        ok("negative: a [trash] part fails 'package'", s["package"] == "FAIL", s["package"])

        t10 = os.path.join(neg, "double_v_repair.xlsx")
        with zipfile.ZipFile(model) as z:
            feed_part = cached.sheet_part(z, "Feed")
        cell_rx = re.compile(r'(<c r="%s"[^>]*>(?:(?!</c>).)*?)(<v>[^<]*</v>)' % addr, re.S)
        zip_patch(model, t10, feed_part, lambda x: cell_rx.sub(r"\1\2\2", x, count=1))
        code, out = xlbridge.check_open(t10, session)
        ok("negative: a cell with two <v> elements fails check-open (exit 4)", code == 4,
           f"exit {code}: {out.splitlines()[-1] if out else ''}")

        # ---- 5. CLI paths
        code = pipeline.main(["verify", run_dir, "--no-open-check", "--session", session])
        ok("cli: verify RUN_DIR exits 0 on the delivered file", code == 0, f"exit {code}")

        cli = os.path.join(test, "cli")
        os.makedirs(cli)
        fresh = os.path.join(cli, "model_under_test.xlsx")
        shutil.copyfile(src, fresh)
        spec2 = os.path.join(cli, "spec.json")
        shutil.copyfile(os.path.join(HERE, "examples", "sample_spec.json"), spec2)
        vj = os.path.join(cli, "values.json")
        code = pipeline.main(["compute", spec2, "-o", vj, "--reader", "stream"])
        got = json.load(open(vj, encoding="utf-8")) if os.path.exists(vj) else {}
        want = json.load(open(os.path.join(run_dir, "inject_values.json"), encoding="utf-8"))
        ok("cli: compute -o (stream reader) writes the same values the run injected",
           code == 0 and got.keys() == want.keys() and all(X.same_value(got[k], want[k]) for k in want),
           f"exit {code}, {len(got)} values")

        owner = os.path.join(cli, "~$model_under_test.xlsx")
        open(owner, "w").close()
        before = sha(fresh)
        code = pipeline.main(["run", spec2, "--run-dir", os.path.join(cli, "run_locked"), "--session", session])
        ok("cli: run refuses a workbook that is open elsewhere (exit 5), file untouched",
           code == 5 and sha(fresh) == before, f"exit {code}")
        os.remove(owner)

        def variant(name, edit):
            sp = json.load(open(spec2, encoding="utf-8"))
            edit(sp)
            path = os.path.join(cli, name)
            json.dump(sp, open(path, "w", encoding="utf-8"), indent=1)
            return path

        bad1 = variant("bad_bare_unit.json", lambda sp: sp["blocks"][0]["periods"].update(label="EUR m"))
        code = pipeline.main(["plan", bad1, "--session", session])
        ok("cli: a column label without its period is a plan error (exit 2)", code == 2, f"exit {code}")
        bad2 = variant("bad_check_label.json",
                       lambda sp: sp["blocks"][0]["rows"][-1].update(label="Check customers"))
        code = pipeline.main(["plan", bad2, "--session", session])
        ok("cli: a check row not labelled 'Check (=0)' is a plan error (exit 2)", code == 2, f"exit {code}")
        bad3 = variant("bad_general.json", lambda sp: sp["blocks"][0].update(format="General"))
        code = pipeline.main(["plan", bad3, "--session", session])
        ok("cli: number format 'General' is refused (exit 2)", code == 2, f"exit {code}")

        merged = os.path.join(cli, "merged_donor.xlsx")
        openpyxl_edit(fresh, merged, lambda ws: ws.merge_cells("B3:C3"))
        shutil.copyfile(merged, fresh)
        before = sha(fresh)
        rd = os.path.join(cli, "run_merged")
        code = pipeline.main(["run", spec2, "--run-dir", rd, "--session", session, "--reader", "stream"])
        failure = json.load(open(os.path.join(rd, "report.json"), encoding="utf-8")).get("failure", "")
        ok("cli: a build failure (merged donor row) exits 4, nothing saved, workbook byte-identical",
           code == 4 and sha(fresh) == before and "merged" in failure, f"exit {code}: {failure[:120]}")

        shutil.copyfile(src, fresh)
        before = sha(fresh)
        real_inject = pipeline.inject

        def boom(*a, **k):
            raise RuntimeError("simulated inject failure after a saved build")
        pipeline.inject = boom
        try:
            rd2 = os.path.join(cli, "run_inject_fails")
            code = pipeline.main(["run", spec2, "--run-dir", rd2, "--session", session, "--reader", "stream"])
        finally:
            pipeline.inject = real_inject
        failed_copy = os.path.join(rd2, "failed_model_under_test.xlsx")
        ok("cli: a failure after the save exits 4, restores the backup byte for byte, keeps the failed file",
           code == 4 and sha(fresh) == before and os.path.exists(failed_copy) and sha(failed_copy) != before,
           f"exit {code}")
    finally:
        print("xl stop:", xlbridge.stop(session), flush=True)

    passed = sum(r["pass"] for r in RESULTS)
    summary = {"passed": passed, "failed": len(RESULTS) - passed, "seconds": round(time.time() - t_start, 1),
               "results": RESULTS}
    with open(os.path.join(test, "selftest_result.json"), "w", encoding="utf-8") as fh:
        json.dump(summary, fh, indent=1)
    print(f"\nSELFTEST {'OK' if passed == len(RESULTS) else 'FAILED'}: {passed}/{len(RESULTS)} passed "
          f"in {summary['seconds']} s")
    return 0 if passed == len(RESULTS) else 1


if __name__ == "__main__":
    sys.exit(main())

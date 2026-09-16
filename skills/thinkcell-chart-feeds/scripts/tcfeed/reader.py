"""Cached values, read without Excel.

Two back-ends for the model values a plan needs:
  stream - openpyxl read-only, streaming: memory stays flat on a 40+ MB model.
  xl     - the xl kernel's cached workbook, `wb(P, values=True)` (skill xl-repl): fast when the
           model is already loaded in that session, but a full load of a large model costs
           gigabytes of kernel memory.
`auto` takes xl below options.reader_mb_limit and stream above it.

Values come back as None (blank), numbers, bools, text, or xlsem.XlError; date-formatted
numbers are returned as Excel serials, which is what a formula sees.
"""
import os
from collections import defaultdict

from . import xlbridge
from . import xlsem as X


def _to_serial(v, epoch=None):
    import datetime
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time, datetime.timedelta)):
        from openpyxl.utils.datetime import to_excel
        return float(to_excel(v, epoch) if epoch is not None else to_excel(v))
    return v


def _conv(cell, epoch=None):
    v = cell.value
    if getattr(cell, "data_type", None) == "e":
        return X.XlError(str(v))
    if not isinstance(v, (str, int, float, bool)) and v is not None:
        v = _to_serial(v, epoch)
    return v


def _sheet(wb, name):
    if name in wb.sheetnames:
        return wb[name]
    low = [s for s in wb.sheetnames if s.lower() == name.lower()]
    if len(low) == 1:
        return wb[low[0]]
    raise KeyError(f"sheet {name!r} not in workbook ({wb.sheetnames})")


class ModelValues:
    def __init__(self, data, backend):
        self.backend = backend
        self._d = {}
        for (s, r, c), v in data.items():
            self._d[(s.lower(), r, c)] = v

    def get(self, sheet, r, c):
        return self._d.get((sheet.lower(), r, c))

    def __len__(self):
        return len(self._d)

    def as_dict(self):
        return dict(self._d)


def _group(needs):
    by = defaultdict(list)
    for s, r0, r1, c0, c1 in needs:
        by[s].append((r0, r1, c0, c1))
    return by


def fetch_stream(path, needs):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    out = {}
    try:
        epoch = getattr(wb, "epoch", None)
        for s, rngs in _group(needs).items():
            ws = _sheet(wb, s)
            r0, r1 = min(r[0] for r in rngs), max(r[1] for r in rngs)
            c0, c1 = min(r[2] for r in rngs), max(r[3] for r in rngs)
            for row in ws.iter_rows(min_row=r0, max_row=r1, min_col=c0, max_col=c1):
                for cell in row:
                    if cell.value is None:
                        continue
                    r, c = cell.row, cell.column
                    if any(a <= r <= b and x <= c <= y for a, b, x, y in rngs):
                        out[(s, r, c)] = _conv(cell, epoch)
    finally:
        wb.close()
    return ModelValues(out, "stream")


_KERNEL_CODE = r'''
import json as _j
from openpyxl.utils.datetime import to_excel as _tx
_W = wb(%(path)r, values=True)
def _cv(cell):
    v = cell.value
    if cell.data_type == "e":
        return {"e": str(v)}
    if isinstance(v, (str, int, float, bool)):
        return v
    return float(_tx(v, getattr(_W, "epoch", None))) if getattr(_W, "epoch", None) is not None else float(_tx(v))
def _sh(name):
    if name in _W.sheetnames:
        return _W[name]
    return _W[[n for n in _W.sheetnames if n.lower() == name.lower()][0]]
_out = []
for _s, _rngs in %(groups)r.items():
    _ws = _sh(_s)
    for (_r, _c), _cell in list(_ws._cells.items()):
        if _cell.value is None:
            continue
        if any(a <= _r <= b and x <= _c <= y for a, b, x, y in _rngs):
            _out.append([_s, _r, _c, _cv(_cell)])
print(_j.dumps(_out))
'''


def fetch_xl(path, needs, session):
    code = _KERNEL_CODE % {"path": os.path.abspath(path).replace("\\", "/"),
                           "groups": {s: list(r) for s, r in _group(needs).items()}}
    resp = xlbridge.exec_code(code, session)
    if not resp.get("ok"):
        raise RuntimeError(f"xl kernel read failed: {resp.get('error') or resp}")
    lines = [ln for ln in (resp.get("stdout") or "").splitlines() if ln.startswith("[")]
    if not lines:
        raise RuntimeError("xl kernel returned no data")
    out = {}
    for s, r, c, v in __import__("json").loads(lines[-1]):
        out[(s, r, c)] = X.XlError(v["e"]) if isinstance(v, dict) else v
    return ModelValues(out, f"xl:{session}")


def fetch(path, needs, backend, session, mb_limit=8):
    if backend == "auto":
        backend = "xl" if os.path.getsize(path) < mb_limit * 1e6 else "stream"
    if backend == "xl":
        return fetch_xl(path, needs, session)
    return fetch_stream(path, needs)


def _all_rows(ws):
    ws.reset_dimensions()          # a stale <dimension> would otherwise cut the stream short
    return ws.iter_rows()


def _ftext(v):
    t = v if isinstance(v, str) else str(getattr(v, "text", v))   # ArrayFormula carries .text
    return t if t.startswith("=") else "=" + t


def scan_sheet(path, sheet):
    """(values, formulas) of every populated cell of one sheet: {(r, c): value} and
    {(r, c): formula text} (formula cells only). Streaming, no Excel."""
    import openpyxl
    values, formulas = {}, {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        epoch = getattr(wb, "epoch", None)
        for row in _all_rows(_sheet(wb, sheet)):
            for cell in row:
                if cell.value is not None:
                    values[(cell.row, cell.column)] = _conv(cell, epoch)
    finally:
        wb.close()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        for row in _all_rows(_sheet(wb, sheet)):
            for cell in row:
                if cell.value is not None and cell.data_type == "f":
                    formulas[(cell.row, cell.column)] = _ftext(cell.value)
    finally:
        wb.close()
    return values, formulas


def compare_workbooks(path_a, path_b, skip_sheets=(), limit=50):
    """Every cached value on every sheet except `skip_sheets`, a against b, row by row
    (memory stays at one row). Returns (cells_compared, differences[:limit], n_diff)."""
    import openpyxl
    wa = openpyxl.load_workbook(path_a, read_only=True, data_only=True, keep_links=False)
    wb_ = openpyxl.load_workbook(path_b, read_only=True, data_only=True, keep_links=False)
    skip = {s.lower() for s in skip_sheets}
    n, diffs, nd = 0, [], 0
    try:
        if [s for s in wa.sheetnames] != [s for s in wb_.sheetnames]:
            diffs.append(("<sheets>", wa.sheetnames, wb_.sheetnames))
            nd += 1
        for name in wa.sheetnames:
            if name.lower() in skip or name not in wb_.sheetnames:
                continue

            def rows(ws, ep):
                for row in _all_rows(ws):
                    d = {}
                    rr = None
                    for cell in row:
                        if cell.value is not None:
                            rr = cell.row
                            d[cell.column] = _conv(cell, ep)
                    if d:
                        yield rr, d
            ia = rows(wa[name], getattr(wa, "epoch", None))
            ib = rows(wb_[name], getattr(wb_, "epoch", None))
            ra, rb = next(ia, None), next(ib, None)
            while ra is not None or rb is not None:
                if rb is None or (ra is not None and ra[0] < rb[0]):
                    cur, da, db = ra[0], ra[1], {}
                    ra = next(ia, None)
                elif ra is None or rb[0] < ra[0]:
                    cur, da, db = rb[0], {}, rb[1]
                    rb = next(ib, None)
                else:
                    cur, da, db = ra[0], ra[1], rb[1]
                    ra, rb = next(ia, None), next(ib, None)
                for c in set(da) | set(db):
                    n += 1
                    if not X.same_value(da.get(c), db.get(c)):
                        nd += 1
                        if len(diffs) < limit:
                            diffs.append((name, cur, c, da.get(c), db.get(c)))
    finally:
        wa.close()
        wb_.close()
    return n, diffs, nd

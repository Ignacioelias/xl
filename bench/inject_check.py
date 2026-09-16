"""End-to-end proof of cached-value injection (think-cell feeds). Install check 6.
Run:  py -3.14 inject_check.py <work_dir>

Builds a workbook whose sheet XML carries every trap the injector has to survive (a self-closing
<v/> with t="str", a shared formula whose master cell has t="e", a formula cell with no <v> at
all, a self-closing non-formula cell sitting right before a formula cell, text and boolean
results, fullCalcOnLoad on). Injects sentinel values that a recalculation would NOT produce,
then opens the file in xl's hidden Excel in manual mode and checks that:
  - Excel shows the sentinels (it reads the injected cached values; xl's instance is in manual
    mode, so recalculation-on-open is guarded by the calcPr warning, not by this check),
  - the workbook did not come back read-only (a repaired package does),
  - calculating the sheet in memory gives the real results (the formulas, shared one included,
    are intact and live),
  - check_open (alerts ON) raises no repair prompt.
Exit 0 pass, 1 fail."""
import os
import re
import shutil
import sys
import zipfile

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font  # noqa: E402

from xl import cached, excelcom  # noqa: E402

SHEET_DATA = (
    '<sheetData>'
    '<row r="1">'
    '<c r="A1"><v>10</v></c>'
    '<c r="B1" t="str"><f>IF(A1="","",A1*2)</f><v/></c>'
    '<c r="C1" s="1"/>'
    '<c r="D1"><f>A1+1</f><v>0</v></c>'
    '<c r="E1" t="str"><f>"FY"&amp;A1</f><v></v></c>'
    '<c r="F1"><f>A1&gt;5</f></c>'
    '</row>'
    '<row r="2"><c r="A2"><v>20</v></c>'
    '<c r="B2" t="e"><f t="shared" ref="B2:B3" si="0">A2*2</f><v>#N/A</v></c></row>'
    '<row r="3"><c r="A3"><v>30</v></c><c r="B3"><f t="shared" si="0"/></c></row>'
    '</sheetData>'
)
SENTINELS = {"B1": 1111.5, "B2": 2222.25, "B3": 3333.125, "C1": 1, "D1": 4444,
             "E1": "FEED", "F1": False, "Z99": 5}
REAL = {"B1": 20, "B2": 40, "B3": 60, "D1": 11, "E1": "FY10", "F1": True}
PATCHED = ["B1", "D1", "E1", "F1", "B2", "B3"]


def build(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Feed"
    ws["C1"].font = Font(bold=True)  # a styled empty cell: openpyxl needs cellXfs index 1 to exist
    ws["A1"] = 10
    wb.save(path)
    tmp = path + ".build"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = data.decode("utf-8")
                text = re.sub(r"<sheetData\s*/>|<sheetData>.*?</sheetData>", SHEET_DATA, text, flags=re.S)
                text = re.sub(r'<dimension ref="[^"]*"\s*/>', '<dimension ref="A1:F3"/>', text)
                data = text.encode("utf-8")
            zout.writestr(item, data)
    os.replace(tmp, path)
    return path


def main(work):
    os.makedirs(work, exist_ok=True)
    src = build(os.path.join(work, "inject_fixture.xlsx"))
    out = os.path.join(work, "inject_fixture_patched.xlsx")
    problems = []

    res = cached.inject(src, "feed", SENTINELS, out=out)  # sheet name resolved case-insensitively
    if res["patched"] != PATCHED:
        problems.append(f"patched {res['patched']} != {PATCHED}")
    if res["no_formula"] != ["C1"] or res["missing"] != ["Z99"]:
        problems.append(f"no_formula {res['no_formula']} / missing {res['missing']}")
    if not (res["warning"] or "").startswith("fullCalcOnLoad"):
        problems.append(f"expected a fullCalcOnLoad warning, got {res['warning']!r}")
    fix = cached.calcpr(out, set={"calcMode": "manual", "fullCalcOnLoad": None}, force=True)
    if fix["calcpr"].get("calcMode") != "manual" or "fullCalcOnLoad" in fix["calcpr"] or fix["warning"]:
        problems.append(f"calcpr not fixed: {fix}")
    if fix.get("backup"):
        os.remove(fix["backup"])
    with zipfile.ZipFile(out) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode()
    if '<f t="shared" ref="B2:B3" si="0">' not in xml or '<f t="shared" si="0"/>' not in xml:
        problems.append("shared formula attributes were altered")
    if '<c r="C1" s="1"/>' not in xml:
        problems.append("the non-formula cell C1 was altered")

    # xl's instance runs in manual mode, so this proves Excel READS the injected values; it cannot
    # prove the file is safe from a recalculation on open in automatic mode. The calcPr warning
    # above is the guard for that (negative control, 2026-09-16: fullCalcOnLoad left on still
    # showed the sentinels here).
    excelcom.app()
    wb = None
    try:
        wb = excelcom._open(out, read_only=False)
        ws = excelcom._ws(wb, "Feed")
        if wb.ReadOnly:
            problems.append("Excel opened the file read-only (package repaired)")
        for ref in PATCHED:
            if ws.Range(ref).Value2 != SENTINELS[ref]:
                problems.append(f"{ref}: Excel shows {ws.Range(ref).Value2!r}, injected {SENTINELS[ref]!r}")
        if ws.Range("B3").Formula != "=A3*2":
            problems.append(f"B3 formula is {ws.Range('B3').Formula!r}")
        ws.Calculate()  # in memory only; the file is closed without saving
        for ref, want in REAL.items():
            if ws.Range(ref).Value2 != want:
                problems.append(f"{ref}: after calculation {ws.Range(ref).Value2!r}, expected {want!r}")
    except Exception as e:  # noqa: BLE001 - a package Excel refuses surfaces here
        problems.append(f"Excel could not open or read the patched file: {e!r}"[:300])
    finally:
        if wb is not None:
            wb.Close(SaveChanges=False)
        excelcom.quit_app()

    chk = excelcom.check_open(out, timeout=90)
    if not chk.get("opened") or chk.get("blocked"):
        problems.append(f"check_open: {chk.get('msg')}")

    for p in problems:
        print("FAIL", p)
    print("inject check:", "ok" if not problems else f"{len(problems)} problem(s)")
    shutil.rmtree(os.path.join(work, "__pycache__"), ignore_errors=True)
    return 0 if not problems else 1


def run(work):
    """main() with any exception reported as a FAIL line (the installer shows the tail)."""
    try:
        return main(work)
    except Exception as e:  # noqa: BLE001
        print("FAIL", repr(e)[:300])
        print("inject check: 1 problem(s)")
        return 1


if __name__ == "__main__":
    sys.exit(run(sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.getcwd(), "inject_check")))

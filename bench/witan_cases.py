"""Five small defect workbooks modelled on the public lint cases in Witan Labs' research
(unsorted approximate lookup, duplicate lookup keys, empty-cell coercion, SUM over text,
mixed currencies), linted with xl. Prints which xl rule fires on each.

The two lookup rules are switched off by default (xl.lint.OFF_IDS), so cases 1 and 2 are
expected to stay silent.

Usage: py -3.14 bench/witan_cases.py
"""
import json
import os
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook

XLCLI = str(Path(__file__).resolve().parent.parent / "xlcli.py")
OUT = Path(os.environ.get("LOCALAPPDATA") or Path.home()) / "xl" / "work" / "witan_cases"
OUT.mkdir(parents=True, exist_ok=True)


def _book(rows, extra):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for r in rows:
        ws.append(r)
    extra(ws)
    return wb


def unsorted_lookup(path):
    def f(ws):
        ws["D2"] = "=VLOOKUP(20,A2:B4,2,TRUE)"
    _book([["Band", "Rate"], [10, 0.1], [30, 0.3], [20, 0.2]], f).save(path)


def duplicate_keys(path):
    def f(ws):
        ws["D2"] = '=VLOOKUP("K1",A2:B4,2,FALSE)'
    _book([["Key", "Name"], ["K1", "alpha"], ["K2", "beta"], ["K1", "gamma"]], f).save(path)


def empty_coercion(path):
    def f(ws):
        ws["A1"] = "=C1*2"
    _book([], f).save(path)


def sum_over_text(path):
    def f(ws):
        ws["A1"], ws["A2"], ws["A3"] = 5, "n/a", 7
        ws["C1"] = "=SUM(A1:A3)"
    _book([], f).save(path)


def mixed_currency(path):
    def f(ws):
        ws["A1"], ws["A2"] = 100, 200
        ws["A1"].number_format = "$#,##0"
        ws["A2"].number_format = "€#,##0"
        ws["A3"] = "=A1+A2"
    _book([], f).save(path)


CASES = [
    (1, "Approximate lookup over an unsorted key range", unsorted_lookup),
    (2, "Exact lookup over duplicate keys", duplicate_keys),
    (3, "Empty cell read as zero in arithmetic", empty_coercion),
    (4, "SUM over a range that holds text", sum_over_text),
    (5, "Two currencies in one formula", mixed_currency),
]


def lint(path):
    r = subprocess.run([sys.executable, XLCLI, "-s", "lintcases", "--json", "lint", str(path)],
                       capture_output=True, text=True)
    try:
        return json.loads(r.stdout)["summary"]["by_rule"]
    except Exception:  # noqa: BLE001
        return {"error": (r.stdout + r.stderr).strip()[-300:]}


if __name__ == "__main__":
    for n, title, make in CASES:
        p = OUT / f"case{n}.xlsx"
        make(p)
        print(f"{n} {title}: {lint(p) or 'no finding'}")
    subprocess.run([sys.executable, XLCLI, "-s", "lintcases", "stop"], capture_output=True)
